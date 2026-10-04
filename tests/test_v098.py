"""v0.9.8: faster envelopes, split piers, stiffness inputs, examples."""

import json

import numpy as np
import pytest

from quickerbridge import engine
from quickerbridge.models import (
    CompositeSlab,
    Model,
    Section,
    concrete_modulus,
    default_model,
)
from quickerbridge.presets import presets
from quickerbridge.projects import SCHEMA_VERSION, create_project, validate_project


def six_spans(split=None, **live):
    m = default_model().model_dump()
    lengths = [18, 35, 45, 35, 18, 18]
    m["spans"] = [
        {"length": length, "section": 0, "zones": [], "simple": i == 0}
        for i, length in enumerate(lengths)
    ]
    m["nonprismatic"] = False
    m["supports"] = ["roller"] * 7
    if split is not None:
        m["supports"][split] = "split"
    m["live"].update(live)
    return Model.model_validate(m)


# --- Exact axle-subset shortcut -------------------------------------------------


@pytest.mark.parametrize(
    "live",
    [
        {"vehicle": "CL750QC", "case": "governing"},
        {"vehicle": "CL625", "case": "truck"},
        {"vehicle": "CL750QC", "dynamic": False},
        {
            "vehicle": "custom",
            "weights": [60, 120, 90],
            "spacings": [3, 5],
            "case": "governing",
        },
    ],
)
def test_truck_selector_equals_every_axle_group(monkeypatch, live):
    m = default_model()
    m.live = m.live.model_copy(update=live)
    m = Model.model_validate(m.model_dump())
    fast = engine.analyse(m)
    monkeypatch.setattr(engine, "truck_selector", lambda groups, count: None)
    slow = engine.analyse(m)
    for sense in ("min", "max"):
        for key in ("V", "M", "D", "R", "Mr"):
            assert np.allclose(fast[sense][key], slow[sense][key], atol=1e-9)
    assert fast["case_max"] == slow["case_max"]
    assert fast["case_min"] == slow["case_min"]


def test_selector_needs_every_subset():
    groups = engine.axle_groups(default_model().live)
    assert engine.truck_selector(groups, 5) is not None
    assert engine.truck_selector(groups[:-1], 5) is None
    special = engine.truck_selector(groups, 5)["special"]
    assert [g["axles"] for g in special] == [[1, 2, 3]]


def test_envelope_close_to_dense_crossings():
    # Crossings are evaluated for their own station only (v0.9.8); the
    # envelopes stay within 0.02 % of the extremes of the 0.9.7 traverse.
    r = engine.analyse(default_model())
    assert max(r["max"]["M"]) == pytest.approx(5675.4, rel=2e-4)
    assert min(r["min"]["M"]) == pytest.approx(-6405.9, rel=2e-4)


def test_crossing_columns_match_full_unit():
    basis = engine.cached_basis(engine.structure_key(default_model()))
    p, cols = engine.crossing_rows(
        basis, np.array([0.0, 3.6, 4.8]), -1, 0.0, basis.length + 4.8
    )
    part = basis.unit_columns(p[:300], cols[:300])
    full = basis.unit(p[:300])
    for slot in range(5):
        c = cols[:300, slot]
        ok = c >= 0
        assert np.allclose(part[ok, slot], full[np.flatnonzero(ok), c[ok]])


def test_lite_spline_gives_the_same_envelope(monkeypatch):
    from quickerbridge import _scipy_lite

    m = six_spans(split=4)
    reference = engine.analyse(m)
    engine.cached_basis.cache_clear()
    monkeypatch.setattr(engine, "CubicSpline", _scipy_lite.CubicSpline)
    lite = engine.analyse(m)
    engine.cached_basis.cache_clear()
    for key in ("V", "M", "R"):
        assert np.allclose(lite["max"][key], reference["max"][key], atol=1e-6)


# --- Split pier ---------------------------------------------------------------------


def test_split_pier_makes_both_sides_independent():
    m = six_spans(split=4)
    basis = engine.cached_basis(engine.structure_key(m))
    x5 = basis.support_x[4]
    # A load in span 5 is carried by supports 5-7 only, as a 2-span beam.
    r = basis.unit([x5 + 9.0])[0][3 * basis.nx : 3 * basis.nx + basis.ns]
    assert np.allclose(r[:4], 0, atol=1e-9)
    assert r[4] == pytest.approx(0.40625) and r[5] == pytest.approx(0.6875)
    # Hinged on both sides: no moment at the split pier.
    result = engine.analyse(m)
    split = result["reactions"][4]["split"]
    assert result["reactions"][4]["type"] == "split"
    assert abs(result["max"]["M"][split["left"]]) < 1e-6
    assert abs(result["min"]["M"][split["right"]]) < 1e-6
    assert result["reactions"][4]["fixity"] == 0


def test_split_pier_reactions_from_the_shear_jump():
    m = six_spans(split=4)
    m.load_mode = "dead"
    result = engine.analyse(m)
    r = result["reactions"][4]
    left = -result["max"]["V"][r["split"]["left"]]
    right = result["max"]["V"][r["split"]["right"]]
    assert left + right == pytest.approx(r["max"])
    assert left > 0 and right > 0


def test_split_pier_only_inside_and_ends_mtq_chain():
    with pytest.raises(ValueError):
        six_spans(split=0)
    m = six_spans(split=4, vehicle="CL750QC", mtq_auto=True)
    basis = engine.cached_basis(engine.structure_key(m))
    low, high = engine.mtq_lane_fractions(m, basis)
    assert low[3 * basis.nx + 4] == high[3 * basis.nx + 4] == 0.63  # no continuity
    from quickerbridge.distribution import effective_spans

    positive, negative = effective_spans(m, 0, 0)
    assert 4 not in [k for k, *_ in negative]  # no M− zone at the split pier


def test_split_pier_in_thermal_modal_and_excel():
    m = six_spans(split=4)
    m.load_mode = "thermal"
    thermal = engine.analyse(m)
    assert thermal["reactions"][4]["split"]
    from quickerbridge.modal import analyse_modal

    assert analyse_modal(m)["modes"]
    pytest.importorskip("openpyxl")
    from io import BytesIO

    from openpyxl import load_workbook
    from quickerbridge.exports import excel_bytes

    m.load_mode = "both"
    wb = load_workbook(BytesIO(excel_bytes(engine.analyse(m), "fr")))
    rows = [r[0] for r in wb["Réactions"].iter_rows(values_only=True)]
    assert "5 gauche" in rows and "5 droite" in rows


# --- Stiffness inputs ---------------------------------------------------------------


def test_direct_ei_from_concrete_or_modulus():
    s = Section(kind="ei", stiffness_input="concrete", I_direct=0.5)
    assert s.E == pytest.approx(concrete_modulus(35, 24) / 1000)
    assert s.EI == pytest.approx(s.E * 1e6 * 0.5)
    s = Section(kind="ei", stiffness_input="modulus", E=30, I_direct=0.2)
    assert s.EI == pytest.approx(30e6 * 0.2)
    assert Section(kind="ei", EI=1234).EI == 1234


def test_nebt_self_weight_follows_unit_weight():
    from quickerbridge.sections import NEBT, properties

    s = Section(kind="nebt", nebt="NEBT1400", unit_weight=25)
    assert properties(s)["w"] == pytest.approx(NEBT["NEBT1400"]["w"] * 25 / 24.5)


def test_inertia_from_composite_is_applied_once():
    from quickerbridge.section_props import section_properties

    slab = CompositeSlab()
    s = Section(composite=slab, inertia_source="1n", inertia_modifier=7)
    p = section_properties(s)
    assert s.inertia_modifier == pytest.approx(
        p["composite"]["1n"]["I"] / p["steel"]["Ix"], rel=1e-6
    )
    # EI = E · I(1n): the girder inertia times M, never M twice.
    from quickerbridge.sections import properties

    assert properties(s)["EI"] == pytest.approx(
        s.E * 1e6 * p["composite"]["1n"]["I"] * 1e-12, rel=1e-6
    )
    manual = Section(composite=slab, inertia_modifier=2.5)
    assert manual.inertia_modifier == 2.5


# --- Examples and projects ------------------------------------------------------------


def test_examples_are_valid_anonymous_and_without_ft():
    items = presets()
    assert len(items) == 4
    for item in items:
        model = Model.model_validate(item["model"])
        assert not model.distribution.enabled and not model.distribution.apply
        text = json.dumps(item, ensure_ascii=False).lower()
        for name in ("yamaska", "melocheville", "hyacinthe"):
            assert name not in text
        project = create_project(model, item["name"]["fr"])
        assert validate_project(json.dumps(project))["schema_version"] == SCHEMA_VERSION
    integral = Model.model_validate(items[2]["model"])
    assert engine.support_fixity(integral)[0] == pytest.approx(0.75, abs=0.01)


def test_schema_11_nebt_project_keeps_its_modulus():
    project = create_project(default_model(), "old")
    project["schema_version"] = 11
    section = project["model"]["sections"][0]
    for key in ("stiffness_input", "I_direct", "fc", "unit_weight", "inertia_source"):
        section.pop(key)
    section.update(kind="nebt", E=28)
    project["model"]["nonprismatic"] = False
    for span in project["model"]["spans"]:
        span["zones"] = []
    opened = validate_project(json.dumps(project))
    assert opened["model"]["sections"][0]["E"] == 28
    assert opened["model"]["sections"][0]["stiffness_input"] == "modulus"
