"""v0.6: rotational-spring supports (partially fixed integral abutments)."""

from io import BytesIO
import json

import numpy as np
from openpyxl import load_workbook
import pytest

from quickerbridge.engine import analyse, cached_basis, snapshot, structure_key
from quickerbridge.exports import excel_bytes
from quickerbridge.models import DeadLoad, Model, Section, Span, ThermalLoad
from quickerbridge.projects import create_project, validate_project
from quickerbridge.thermal import analyse_thermal

EI, L, W = 2.0e6, 20.0, 10.0


def spring_beam(supports, springs, spans=(L,), mode="dead"):
    return Model(
        spans=[Span(length=s) for s in spans],
        supports=supports,
        support_springs=springs,
        sections=[Section(kind="ei", EI=EI)],
        dead=[DeadLoad(w=W)],
        load_mode=mode,
    )


@pytest.mark.parametrize("k", [1e3, 3e4, 3e5, 1e6, 1e8])
def test_spring_propped_beam_matches_compatibility(k):
    # Rotation of a simple span θ0 = wL³/24EI is shared by the spring and the
    # end moment: M = θ0 / (L/3EI + 1/k).
    r = analyse(spring_beam(["spring", "pin"], [k, 0]))
    expected = W * L**3 / (24 * EI) / (L / (3 * EI) + 1 / k)
    assert r["max"]["M"][0] == pytest.approx(-expected, rel=1e-6)
    assert r["reactions"][0]["moment_max"] == pytest.approx(expected, rel=1e-6)
    assert r["reactions"][0]["fixity"] == pytest.approx(k / (k + 3 * EI / L))
    assert r["reactions"][1]["fixity"] == 0


@pytest.mark.parametrize("k", [5e4, 5e5, 5e6])
def test_symmetric_springs_both_ends(k):
    r = analyse(spring_beam(["spring", "spring"], [k, k]))
    expected = W * L**3 / (24 * EI) / (L / (2 * EI) + 1 / k)
    assert r["max"]["M"][0] == pytest.approx(-expected, rel=1e-6)
    assert r["max"]["M"][-1] == pytest.approx(-expected, rel=1e-6)
    assert r["reactions"][0]["moment_max"] == pytest.approx(
        -r["reactions"][1]["moment_max"], rel=1e-9
    )


def test_spring_limits_bracket_pinned_and_fixed():
    pinned = analyse(spring_beam(["pin", "roller"], []))
    fixed = analyse(spring_beam(["fixed", "fixed"], []))
    soft = analyse(spring_beam(["spring", "spring"], [1e-3, 1e-3]))
    stiff = analyse(spring_beam(["spring", "spring"], [1e13, 1e13]))
    for key in ("M", "D"):
        np.testing.assert_allclose(
            soft["max"][key], pinned["max"][key], atol=1e-4 * max(pinned["max"][key])
        )
        np.testing.assert_allclose(
            stiff["max"][key], fixed["max"][key], atol=1e-4 * max(fixed["max"][key])
        )
    mid = analyse(spring_beam(["spring", "spring"], [6e5, 6e5]))
    i = len(mid["x"]) // 2
    assert fixed["max"]["M"][i] < mid["max"]["M"][i] < pinned["max"]["M"][i]


def test_spring_moving_load_equilibrium_and_symmetry():
    model = Model(
        supports=["spring", "pin", "spring"],
        support_springs=[2e6, 0, 2e6],
        load_mode="live",
    )
    basis = cached_basis(structure_key(model))
    p = np.linspace(0, basis.length, 801)
    u = basis.unit(p)
    nx, ns = basis.nx, basis.ns
    r, mr = u[:, 3 * nx : 3 * nx + ns], u[:, 3 * nx + ns :]
    assert np.abs(r.sum(axis=1) - 1).max() < 1e-9
    assert np.abs(r @ basis.support_x + mr.sum(axis=1) - p).max() < 1e-7
    result = analyse(model)
    x = np.array(result["x"])
    for sense in ("min", "max"):
        v = np.array(result[sense]["M"])
        assert np.abs(v - np.interp(x[-1] - x, x, v)).max() < 1e-6 * np.abs(v).max()
    nxr = len(result["x"])
    snap = snapshot(model, result["case_max"][nxr + 5], nxr + 5, "max")
    assert abs(snap["equilibrium"]["moment"]) < 1e-5


def test_interior_rotational_spring():
    model = spring_beam(["pin", "spring", "roller"], [0, 1e6, 0], spans=(15, 25))
    r = analyse(model)
    assert abs(r["reactions"][1]["moment_max"]) > 1e-6
    assert 0 < r["reactions"][1]["fixity"] < 1


def test_thermal_with_springs_between_pinned_and_fixed():
    def run(supports, springs):
        return analyse_thermal(
            Model(
                spans=[Span(length=L)],
                supports=supports,
                support_springs=springs,
                sections=[Section(kind="ei", EI=EI)],
                load_mode="thermal",
                thermal=ThermalLoad(delta_T=20),
            )
        )

    fixed = run(["fixed", "fixed"], [])["values"]["M"][0]
    spring = run(["spring", "spring"], [3e5, 3e5])
    kappa = spring["meta"]["curvature"]
    # Symmetric springs: uniform M with θ = M/k = (−κ·L/2 − M·L/2EI)… closed form:
    expected = -kappa * L / 2 / (L / (2 * EI) + 1 / 3e5)
    np.testing.assert_allclose(spring["values"]["M"], expected, rtol=1e-6)
    assert 0 < expected < fixed
    assert spring["meta"]["equilibrium_error"] < 1e-9


def test_spring_validation_and_project_round_trip():
    with pytest.raises(ValueError):
        spring_beam(["spring", "pin"], [])
    with pytest.raises(ValueError):
        spring_beam(["spring", "pin"], [0, 0])
    model = spring_beam(["spring", "roller"], [4.5e5, 0])
    project = create_project(model, "Spring")
    assert validate_project(json.dumps(project))["model"] == model.model_dump(
        mode="json"
    )


def test_excel_lists_spring_stiffness():
    r = analyse(spring_beam(["spring", "roller"], [4.5e5, 0]))
    wb = load_workbook(BytesIO(excel_bytes(r, "en")))
    sheet = wb["Reactions"]
    header = [c.value for c in sheet[1]]
    kind = sheet.cell(2, header.index("Type") + 1).value
    assert "spring" in kind and "450000" in kind.replace(",", "")
