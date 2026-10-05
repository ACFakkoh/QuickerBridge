"""v0.9.7: full-deck Canadian lane loads, MTQ automatic fraction, both directions,
up to 7 spans, slab shrinkage and creep as imposed deformations."""

import json

import numpy as np
import pytest

from quickerbridge.browser import dispatch
from quickerbridge.engine import analyse, position_record, snapshot
from quickerbridge.loads import lane_parameters
from quickerbridge.models import CompositeSlab, LiveLoad, Model, Section, Span
from quickerbridge.projects import create_project, validate_project
from quickerbridge.thermal import analyse_thermal, free_curvature, slab_restraint


def reference(vehicle, fraction=0.8, case="governing", mtq_auto=False):
    """2 x 34.8 m, steel alone, prismatic S1 (1200, 350x25 / 600x50), live only."""
    model = Model(sections=[Section()], dead=[], load_mode="live")
    model.self_weight.apply = False
    model.live.vehicle = vehicle
    model.live.lane_fraction = fraction
    model.live.case = case
    model.live.mtq_auto = mtq_auto
    model.live.lane_extent = "full"  # v0.9.7 reference values (whole deck)
    result = analyse(model)
    return max(result["max"]["M"]), min(result["min"]["M"])


def test_lanes_always_cover_the_whole_deck():
    for vehicle in ("CL625", "CL750QC", "custom"):
        assert lane_parameters(LiveLoad(vehicle=vehicle))[2] == "full"


def test_cl750qc_80_lane_case_full_deck():
    m_pos, m_neg = reference("CL750QC", 0.8, "lane")
    assert m_pos == pytest.approx(4002.0, abs=1.0)
    assert m_neg == pytest.approx(-3659.4, abs=1.0)


def test_truck_governs_positive_moment_for_both_fractions():
    for fraction in (0.8, 0.63):
        m_pos, _ = reference("CL750QC", fraction)
        assert m_pos == pytest.approx(4578.6, abs=1.0)  # x1.25 truck
    lane63, lane63_neg = reference("CL750QC", 0.63, "lane")
    assert lane63 == pytest.approx(3379.3, abs=1.0)
    assert lane63_neg == pytest.approx(-3287.1, abs=1.0)


def test_cl625_envelope():
    m_pos, m_neg = reference("CL625")
    assert m_pos == pytest.approx(3766.6, abs=1.0)
    assert m_neg == pytest.approx(-2818.2, abs=1.0)


def test_direction_is_always_both():
    assert LiveLoad(direction="forward").direction == "both"
    forward = Model(live=LiveLoad(direction="forward"), load_mode="live")
    both = Model(load_mode="live")
    np.testing.assert_allclose(analyse(forward)["max"]["M"], analyse(both)["max"]["M"])
    project = create_project(forward, "dir")
    project["model"]["live"]["direction"] = "reverse"
    reopened = validate_project(json.dumps(project))
    assert reopened["model"]["live"]["direction"] == "both"


def test_seven_spans():
    model = Model(
        spans=[Span(length=20) for _ in range(7)],
        supports=["roller"] + ["pin"] * 6 + ["roller"],
        load_mode="live",
    )
    result = analyse(model)
    assert len(result["reactions"]) == 8
    with pytest.raises(ValueError):
        Model(spans=[Span() for _ in range(8)], supports=["pin"] * 9)


def test_manual_position_uses_full_deck_lane():
    model = Model(load_mode="live", live=LiveLoad(vehicle="CL625", case="lane"))
    record = position_record(model, 20.0, "forward")
    assert record["manual"]
    view = snapshot(model, record)
    assert view["lane"] == [{"start": 0.0, "end": 69.6, "w": 9.0}]
    assert abs(view["equilibrium"]["force"]) < 1e-6


def slab_model(imposed):
    model = Model(sections=[Section(composite=CompositeSlab())], load_mode="thermal")
    model.thermal.imposed = imposed
    return model


def test_shrinkage_curvature_and_signs():
    model = slab_model("shrinkage")
    section = model.sections[0]
    r = slab_restraint(model, section)
    expected = -250e-6 * r["area"] * r["lever"] / r["I"] * 1000
    assert free_curvature(model, section) == pytest.approx(expected)
    assert expected < 0  # sagging: the slab shortens
    result = analyse_thermal(model)
    m = np.array(result["values"]["M"])
    assert m.min() < -100 and m.max() <= 1e-6  # hogging restraint at the pier
    assert max(result["values"]["D"]) > 0  # bows downward
    assert abs(sum(result["values"]["R"])) < 1e-6
    assert result["meta"]["slab_strain"] == pytest.approx(250)
    # Twice the strain, twice the effects (linear).
    model.thermal.shrinkage_micro = 500
    np.testing.assert_allclose(analyse_thermal(model)["values"]["M"], 2 * m, atol=1e-6)


def test_creep_is_phi_sigma_over_ec_and_independent_of_thermal_inputs():
    from quickerbridge.section_props import concrete_modulus

    creep = slab_model("creep")
    ec = concrete_modulus(35, 24 * 1000 / 9.81)
    result = analyse_thermal(creep)
    assert result["meta"]["slab_strain"] == pytest.approx(2.0 * 3.0 / ec * 1e6)
    creep.thermal.delta_T = -40  # thermal inputs do not affect creep
    np.testing.assert_allclose(
        analyse_thermal(creep)["values"]["M"], result["values"]["M"], atol=1e-9
    )


def test_shrinkage_needs_a_slab():
    model = Model(load_mode="thermal")
    model.thermal.imposed = "shrinkage"
    with pytest.raises(ValueError, match="thermal.needs_slab"):
        analyse_thermal(model)
    with pytest.raises(ValueError, match="thermal.needs_slab"):
        dispatch(
            json.dumps(
                {"action": "analyse", "data": {"model": model.model_dump(), "job": "s"}}
            )
        )


def mtq_run(auto, fraction=0.8, spans=2, length=34.8):
    model = Model(
        spans=[Span(length=length) for _ in range(spans)],
        supports=["roller"] + ["pin"] * (spans - 1) + ["roller"],
        sections=[Section()],
        dead=[],
        load_mode="live",
    )
    model.self_weight.apply = False
    model.live.case = "lane"
    model.live.mtq_auto = auto
    model.live.lane_fraction = fraction
    model.live.lane_extent = "full"  # v0.9.7 reference values (whole deck)
    return model, analyse(model)


def test_mtq_automatic_fraction_by_response():
    model, auto = mtq_run(True)
    _, f80 = mtq_run(False, 0.8)
    _, f63 = mtq_run(False, 0.63)
    assert auto["vehicle"]["mtq_auto"] and not f80["vehicle"]["mtq_auto"]
    x = np.array(auto["x"])
    zone = (x >= 27.84 - 1e-6) & (x <= 41.76 + 1e-6)  # 0.20 L either side of the pier
    m_hi = np.array(auto["max"]["M"])
    np.testing.assert_allclose(m_hi[~zone], np.array(f63["max"]["M"])[~zone], atol=1e-6)
    np.testing.assert_allclose(m_hi[zone], np.array(f80["max"]["M"])[zone], atol=1e-6)
    assert min(auto["min"]["M"]) == pytest.approx(-3659.4, abs=1.0)
    np.testing.assert_allclose(auto["min"]["M"], f80["min"]["M"], atol=1e-6)
    for sense in ("min", "max"):
        np.testing.assert_allclose(auto[sense]["V"], f63[sense]["V"], atol=1e-6)
        np.testing.assert_allclose(auto[sense]["D"], f80[sense]["D"], atol=1e-6)
    for sense in ("min", "max"):
        reaction = lambda r, i: r[sense]["R"][i]  # noqa: E731
        assert reaction(auto, 0) == pytest.approx(reaction(f63, 0))
        assert reaction(auto, 2) == pytest.approx(reaction(f63, 2))
        assert reaction(auto, 1) == pytest.approx(reaction(f80, 1))
    # The records carry the fraction used, and a click reproduces it.
    for e in auto["extrema"]:
        view = snapshot(model, e, e["index"], e["sense"])
        assert abs(view["equilibrium"]["force"]) < 1e-6
        assert view["lane"] == [{"start": 0.0, "end": 69.6, "w": 12.6}]
        if e["response"] == "V":
            assert e["fraction"] == 0.63
        if e["response"] == "M" and e["sense"] == "min":
            assert e["fraction"] == 0.8
    assert auto["case_max"][auto["extrema"][2]["index"]]["fraction"] == 0.63


def test_mtq_single_span_and_hinge_use_63():
    _, auto = mtq_run(True, spans=1, length=30)
    _, f63 = mtq_run(False, 0.63, spans=1, length=30)
    for sense in ("min", "max"):
        for key in "VMDR":
            np.testing.assert_allclose(auto[sense][key], f63[sense][key], atol=1e-6)
    model, _ = mtq_run(True, spans=3, length=30)
    model.spans[1].simple = True
    result = analyse(model)
    nx = len(result["x"])
    # Supports next to the simple span have no deck continuity: 63 %.
    for k in (1, 2):
        assert result["case_max"][3 * nx + k]["fraction"] == 0.63


def test_mtq_integral_abutment_zone_keeps_63_for_positive_moment():
    """Only the M− zones over interior piers take 80 % for M+ (user rule)."""
    from quickerbridge.engine import Basis, mtq_lane_fractions

    model, _ = mtq_run(True)
    model.supports = ["fixed", "pin", "fixed"]
    basis = Basis(model)
    _, high = mtq_lane_fractions(model, basis)
    x, nx = basis.x, basis.nx
    m_hi = high[nx : 2 * nx]
    near_abutment = x <= 0.15 * 34.8 - 1e-6
    near_pier = (x >= 27.84 + 1e-6) & (x <= 41.76 - 1e-6)
    assert np.all(m_hi[near_abutment] == 0.63)
    assert np.all(m_hi[near_pier] == 0.8)


def test_mtq_default_on_for_old_projects_and_manual_position():
    project = create_project(Model(), "old")
    del project["model"]["live"]["mtq_auto"]
    assert validate_project(json.dumps(project))["model"]["live"]["mtq_auto"] is True
    model = Model(load_mode="live", live=LiveLoad(case="lane"))
    assert position_record(model, 20.0, "forward")["fraction"] == 0.8


def test_dead_only_exports_one_value_per_effect():
    """Permanent loads alone: no min/max duplication in CSV and Excel."""
    from io import BytesIO

    from openpyxl import load_workbook

    from quickerbridge.exports import csv_bytes, excel_bytes

    model = Model(load_mode="dead")
    result = analyse(model)
    header = csv_bytes(result).decode("utf-8-sig").splitlines()[0]
    assert "V (kN)" in header and "V min" not in header
    wb = load_workbook(BytesIO(excel_bytes(result, model=model)))
    stations = [c.value for c in wb["Stations"][1]]
    assert "M (kN·m)" in stations and "M max (kN·m)" not in stations
    assert "Governing cases" not in wb.sheetnames
    reactions = [c.value for c in wb["Reactions"][1]]
    assert reactions[2] == "R (kN)" and "R max (kN)" not in reactions
    total = sum(row[2].value for row in wb["Reactions"].iter_rows(min_row=2))
    assert total == pytest.approx(sum(r["max"] for r in result["reactions"]))
