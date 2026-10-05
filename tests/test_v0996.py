"""v0.9.96: lane UDL placed where it increases the effect (S6 C3.8.4.1),
S6 pedestrian load over every combination of loaded spans, project schema 13,
bilinear gradient default and Excel metadata."""

import io
import json

import numpy as np
import pytest

from quickerbridge.engine import analyse, snapshot
from quickerbridge.exports import excel_bytes
from quickerbridge.models import Model, Section, Span, ThermalLoad
from quickerbridge.projects import create_project, validate_project


def bridge(lengths, **live):
    model = Model(
        spans=[Span(length=L) for L in lengths],
        supports=["roller"] + ["pin"] * (len(lengths) - 1) + ["roller"],
        sections=[Section()],
        dead=[],
        load_mode="live",
    )
    model.self_weight.apply = False
    for key, value in live.items():
        setattr(model.live, key, value)
    return model


def test_lane_extent_default_and_ordering():
    assert Model().live.lane_extent == "spans"
    full = analyse(
        bridge([30, 40, 30], vehicle="CL625", case="lane", lane_extent="full")
    )
    spans = analyse(bridge([30, 40, 30], vehicle="CL625", case="lane"))
    infl = analyse(
        bridge([30, 40, 30], vehicle="CL625", case="lane", lane_extent="influence")
    )
    for key in ("V", "M", "D", "R"):
        # Placing the UDL only where it helps can never reduce an extreme.
        assert np.all(np.array(spans["max"][key]) >= np.array(full["max"][key]) - 1e-6)
        assert np.all(np.array(spans["min"][key]) <= np.array(full["min"][key]) + 1e-6)
        # Influence parts: numerical areas (piecewise-linear lines), 1e-4.
        tol = 1e-4 * np.abs(spans["max"][key]).max()
        assert np.all(np.array(infl["max"][key]) >= np.array(spans["max"][key]) - tol)
        assert np.all(np.array(infl["min"][key]) <= np.array(spans["min"][key]) + tol)
    assert max(spans["max"]["M"]) > max(full["max"]["M"]) * 1.05


def test_lane_udl_alone_matches_hand_values():
    # Simple span, lane case: the UDL part at midspan is w L^2 / 8 whatever
    # the extent (the influence line has one sign).
    for extent in ("full", "spans", "influence"):
        r = analyse(
            bridge(
                [20],
                vehicle="custom",
                case="lane",
                lane_extent=extent,
                weights=[1e-6],
                spacings=[],
                lane_w=10.0,
                lane_fraction=0.8,
                dynamic=False,
            )
        )
        assert max(r["max"]["M"]) == pytest.approx(10 * 20**2 / 8, rel=2e-3)
    # Two equal spans, UDL alone: M+ with one span loaded 0.0957 w L^2 (exact
    # 0.095703), M- with both spans 0.125 w L^2.
    r = analyse(
        bridge(
            [20, 20],
            vehicle="custom",
            case="lane",
            weights=[1e-6],
            spacings=[],
            lane_w=10.0,
            dynamic=False,
        )
    )
    assert max(r["max"]["M"]) == pytest.approx(0.095703 * 10 * 400, rel=3e-3)
    assert min(r["min"]["M"]) == pytest.approx(-0.125 * 10 * 400, rel=3e-3)


@pytest.mark.parametrize("extent", ["spans", "influence"])
def test_lane_snapshots_reproduce_the_envelope(extent):
    model = bridge([20, 30, 20], vehicle="CL750QC", case="lane", lane_extent=extent)
    r = analyse(model)
    nx = len(r["x"])
    for j in range(0, 3 * nx, 5):
        for sense in ("min", "max"):
            record = r["case_" + sense][j]
            if record["case"] == "unloaded":
                continue
            view = snapshot(model, record, j, sense)
            got = np.r_[view["V"], view["M"], view["D"]][j]
            assert got == pytest.approx(
                r[sense]["V" if j < nx else "M" if j < 2 * nx else "D"][j % nx],
                rel=1e-6 if extent == "spans" else 1e-4,
                abs=1e-6,
            )
            assert abs(view["equilibrium"]["force"]) < 1e-6
            assert abs(view["equilibrium"]["moment"]) < 1e-5


def test_lane_snapshot_loads_only_the_helpful_spans():
    model = bridge([30, 30], vehicle="CL625", case="lane")
    r = analyse(model)
    nx = len(r["x"])
    k = int(np.argmax(r["max"]["M"]))
    view = snapshot(model, r["case_max"][nx + k], nx + k, "max")
    assert view["lane"] == [{"start": 0.0, "end": 30.0, "w": 9.0}] or view["lane"] == [
        {"start": 30.0, "end": 60.0, "w": 9.0}
    ]


def test_pedestrian_load_hand_values():
    # 20 m: s = 20 m, p = min(5 - 20/30, 4) = 4 kPa, 2 m wide: 8 kN/m.
    m = bridge([20], source="pedestrian")
    r = analyse(m)
    assert max(r["max"]["M"]) == pytest.approx(8 * 20**2 / 8, rel=1e-6)
    assert r["vehicle"]["weights"] == [] and r["vehicle"]["pedestrian_width"] == 2.0
    # 2 x 30 m: M+ with one span (p = 4), M- with both (s = 60, p = 3).
    r = analyse(bridge([30, 30], source="pedestrian"))
    assert max(r["max"]["M"]) == pytest.approx(0.095703 * 8 * 900, rel=2e-3)
    assert min(r["min"]["M"]) == pytest.approx(-0.125 * 6 * 900, rel=1e-6)
    e = next(x for x in r["extrema"] if x["response"] == "M" and x["sense"] == "min")
    assert e["case"] == "pedestrian" and e["spans"] == [1, 2]
    assert e["intensity"] == pytest.approx(3.0)


def test_pedestrian_combinations_beat_sign_rule():
    # p depends on s: loading an extra span of the right sign can lower p for
    # the others. Every combination is checked: the envelope is at least the
    # best single-span and all-same-sign arrangements.
    m = bridge([40, 10, 40], source="pedestrian")
    m.pedestrian.p_min = 0.5
    r = analyse(m)
    nx = len(r["x"])
    for j in range(0, 3 * nx, 7):
        for sense in ("min", "max"):
            record = r["case_" + sense][j]
            if record["case"] != "pedestrian":
                continue
            view = snapshot(m, record, j, sense)
            got = np.r_[view["V"], view["M"], view["D"]][j]
            ref = r[sense]["V" if j < nx else "M" if j < 2 * nx else "D"][j % nx]
            assert got == pytest.approx(ref, rel=1e-9, abs=1e-9)


def test_pedestrian_slab_width_and_bounds():
    m = bridge([10], source="pedestrian")
    m.sections[0].composite = __import__(
        "quickerbridge.models", fromlist=["CompositeSlab"]
    ).CompositeSlab(effective_width=3000)
    r = analyse(m)
    assert r["vehicle"]["pedestrian_width"] == pytest.approx(3.0)
    assert max(r["max"]["M"]) == pytest.approx(4 * 3 * 100 / 8, rel=1e-6)
    m.pedestrian.width_source = "manual"
    m.pedestrian.width = 1500
    assert max(analyse(m)["max"]["M"]) == pytest.approx(4 * 1.5 * 100 / 8, rel=1e-6)
    with pytest.raises(ValueError):
        Model.model_validate({"pedestrian": {"p_min": 5, "p_max": 4}})


def test_pedestrian_and_maintenance_are_never_added():
    seen = set()
    for length in (12, 40):  # maintenance governs short spans, pedestrians long
        ped = analyse(bridge([length], source="pedestrian"))
        veh = analyse(bridge([length], vehicle="Maintenance", case="truck"))
        m = bridge([length], source="pedestrian")
        m.pedestrian.maintenance = True
        both = analyse(m)
        for key in ("V", "M", "D"):
            np.testing.assert_allclose(
                both["max"][key],
                np.maximum(ped["max"][key], veh["max"][key]),
                atol=1e-9,
            )
            np.testing.assert_allclose(
                both["min"][key],
                np.minimum(ped["min"][key], veh["min"][key]),
                atol=1e-9,
            )
        seen |= {c["case"] for c in both["case_max"]}
        e = next((x for x in both["extrema"] if x["case"] == "truck"), None)
        if e:  # a maintenance extreme reproduces with the maintenance vehicle
            view = snapshot(m, e, e["index"], e["sense"])
            assert sorted(a["load"] for a in view["axles"]) in ([24.0, 56.0], [56.0])
    assert {"pedestrian", "truck"} <= seen


def test_pedestrian_ignores_ft_and_axle_factor():
    m = bridge([25, 25], source="pedestrian")
    plain = analyse(m)
    m.live.axle_factor = 3.0
    m.distribution.enabled = m.distribution.apply = True
    r = analyse(m)
    np.testing.assert_allclose(r["max"]["M"], plain["max"]["M"])
    assert r["ft"] is None
    m.live.factor = 1.7
    np.testing.assert_allclose(
        analyse(m)["max"]["M"], 1.7 * np.array(plain["max"]["M"]), rtol=1e-9
    )


def test_project_schema_13_round_trip():
    m = Model()
    m.live.source = "pedestrian"
    m.live.lane_extent = "influence"
    m.pedestrian.maintenance = True
    text = json.dumps(create_project(m, "piétons"))
    data = validate_project(text)
    assert data["schema_version"] == 13
    assert data["model"]["live"]["source"] == "pedestrian"
    assert data["model"]["pedestrian"]["maintenance"] is True
    old = json.loads(text)
    old["schema_version"] = 12
    del old["model"]["pedestrian"]
    del old["model"]["live"]["source"]
    del old["model"]["live"]["lane_extent"]
    data = validate_project(json.dumps(old))
    assert data["model"]["live"]["lane_extent"] == "spans"
    assert data["model"]["live"]["source"] == "vehicle"


def test_bilinear_gradient_default_30():
    assert ThermalLoad().slab_delta_T == 30


def test_excel_carries_name_date_and_version():
    from openpyxl import load_workbook
    from quickerbridge.version import APP_VERSION

    m = bridge([20], source="pedestrian")
    r = analyse(m)
    wb = load_workbook(
        io.BytesIO(
            excel_bytes(r, "fr", m, {"name": "Pont X", "exported": "2026-10-04 12:00"})
        )
    )
    rows = [tuple(c.value for c in row) for row in wb["Modèle"].iter_rows()]
    assert ("QuickerBridge", APP_VERSION) in rows
    assert ("Nom du modèle", "Pont X") in rows
    assert ("Date d’export", "2026-10-04 12:00") in rows
    gov = [row for row in wb["Cas déterminants"].iter_rows(values_only=True)]
    assert any(str(row[6]).startswith("spans/travées 1") for row in gov[1:])
