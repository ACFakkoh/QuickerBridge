"""v0.5: integral (fixed) abutments, zone plate source, influence lines, traverse."""

from quickerbridge.projects import SCHEMA_VERSION
from io import BytesIO
import json

import numpy as np
from openpyxl import load_workbook
import pycba as cba
import pytest

from quickerbridge.browser import dispatch
from quickerbridge.engine import analyse, cached_basis, snapshot, structure_key
from quickerbridge.exports import excel_bytes
from quickerbridge.models import (
    DeadLoad,
    LiveLoad,
    Model,
    Section,
    Span,
    ThermalLoad,
    Zone,
)
from quickerbridge.projects import create_project, validate_project
from quickerbridge.sections import span_ei
from quickerbridge.thermal import analyse_thermal

EI = 2.0e6  # kN m²


def beam(spans, supports, mode="dead", w=10.0, **kwargs):
    return Model(
        spans=[Span(length=length) for length in spans],
        supports=supports,
        sections=[Section(kind="ei", EI=EI)],
        dead=[DeadLoad(w=w)],
        load_mode=mode,
        **kwargs,
    )


def station(result, x, side="right"):
    xs = np.array(result["x"])
    candidates = np.flatnonzero(np.isclose(xs, xs[np.abs(xs - x).argmin()]))
    for i in candidates:
        if result["sides"][i] == side:
            return int(i)
    return int(candidates[0])


# ---------------------------------------------------------------- analytics


def test_fixed_fixed_udl_matches_closed_form():
    L, w = 20.0, 10.0
    r = analyse(beam([L], ["fixed", "fixed"]))
    M, D = np.array(r["max"]["M"]), np.array(r["max"]["D"])
    assert M[0] == pytest.approx(-w * L**2 / 12, rel=1e-6)
    assert M[-1] == pytest.approx(-w * L**2 / 12, rel=1e-6)
    assert M[station(r, L / 2)] == pytest.approx(w * L**2 / 24, rel=1e-6)
    assert D[station(r, L / 2)] == pytest.approx(
        1000 * w * L**4 / (384 * EI), rel=1e-4
    )
    left, right = r["reactions"]
    assert left["max"] == pytest.approx(w * L / 2, rel=1e-9)
    assert left["type"] == "fixed"
    # Moment reactions are counter-clockwise positive (PyCBA convention).
    assert left["moment_max"] == pytest.approx(w * L**2 / 12, rel=1e-6)
    assert right["moment_max"] == pytest.approx(-w * L**2 / 12, rel=1e-6)


def test_propped_cantilever_udl():
    L, w = 15.0, 8.0
    r = analyse(beam([L], ["fixed", "roller"], w=w))
    assert r["max"]["M"][0] == pytest.approx(-w * L**2 / 8, rel=1e-6)
    assert r["reactions"][0]["max"] == pytest.approx(5 * w * L / 8, rel=1e-9)
    assert r["reactions"][1]["max"] == pytest.approx(3 * w * L / 8, rel=1e-9)
    assert r["reactions"][1]["moment_max"] == 0


def test_pinned_supports_have_no_moment_reaction():
    r = analyse(beam([10, 12], ["pin", "roller", "roller"]))
    assert all(abs(x["moment_max"]) < 1e-12 for x in r["reactions"])


@pytest.mark.parametrize("a", [2.0, 5.0, 7.5])
def test_fixed_fixed_influence_matches_closed_form(a):
    L = 10.0
    model = beam([L], ["fixed", "fixed"], mode="live")
    basis = cached_basis(structure_key(model))
    u = basis.unit([a])[0]
    b = L - a
    nx, ns = basis.nx, basis.ns
    assert u[nx + 0] == pytest.approx(-a * b**2 / L**2, abs=2e-4)  # M at left end
    assert u[3 * nx + ns] == pytest.approx(a * b**2 / L**2, abs=2e-4)  # Mr CCW +
    assert u[3 * nx] == pytest.approx(b**2 * (L + 2 * a) / L**3, abs=2e-5)


def test_moving_load_equilibrium_includes_moment_reactions():
    model = beam([18, 24, 18], ["fixed", "pin", "roller", "fixed"], mode="live")
    basis = cached_basis(structure_key(model))
    p = np.linspace(0, basis.length, 1201)
    u = basis.unit(p)
    nx, ns = basis.nx, basis.ns
    r, mr = u[:, 3 * nx : 3 * nx + ns], u[:, 3 * nx + ns :]
    assert np.abs(r.sum(axis=1) - 1).max() < 1e-9
    assert np.abs(r @ basis.support_x + mr.sum(axis=1) - p).max() < 1e-7
    # Sagging M at the right fixed end equals its CCW moment reaction.
    assert np.abs(u[:, 2 * nx - 1] - mr[:, -1]).max() < 1e-7


def test_integral_two_span_envelope_is_mirror_symmetric():
    model = Model(
        supports=["fixed", "pin", "fixed"], load_mode="both", precision="standard"
    )
    r = analyse(model)
    x = np.array(r["x"])
    for key in ("M", "D"):
        for sense in ("min", "max"):
            v = np.array(r[sense][key])
            mirrored = np.interp(x[-1] - x, x, v)
            assert np.abs(v - mirrored).max() < 1e-6 * np.abs(v).max() + 1e-9
    left, right = r["reactions"][0], r["reactions"][-1]
    assert left["moment_max"] == pytest.approx(-right["moment_min"], rel=1e-9)
    assert left["moment_max"] > 0


def test_fixing_abutments_reduces_sagging_and_deflection():
    pinned = analyse(Model(load_mode="live"))
    integral = analyse(Model(load_mode="live", supports=["fixed", "pin", "fixed"]))
    assert max(integral["max"]["M"]) < max(pinned["max"]["M"])
    assert max(integral["max"]["D"]) < max(pinned["max"]["D"])


def test_snapshot_equilibrium_with_fixed_abutments():
    model = Model(supports=["fixed", "roller", "fixed"], load_mode="both")
    result = analyse(model)
    nx = len(result["x"])
    for index in (nx + 10, nx + 80, 3 * nx + 1):
        for sense in ("max", "min"):
            snap = snapshot(model, result["case_" + sense][index], index, sense)
            assert abs(snap["equilibrium"]["force"]) < 1e-6
            assert abs(snap["equilibrium"]["moment"]) < 1e-5
            # Plotted diagram closes at the right end: M(L) = Mr(right), CCW +.
            assert snap["plot"]["M"][-1] == pytest.approx(snap["Mr"][-1], abs=1e-6)


def test_thermal_fixed_fixed_gives_uniform_moment_and_no_deflection():
    L = 12.0
    model = Model(
        spans=[Span(length=L)],
        supports=["fixed", "fixed"],
        sections=[Section(kind="ei", EI=EI)],
        load_mode="thermal",
        thermal=ThermalLoad(
            delta_T=20, alpha_micro=12, depth=1500, depth_source="manual"
        ),
    )
    r = analyse_thermal(model)
    kappa = r["meta"]["curvature"]
    np.testing.assert_allclose(r["values"]["M"], -EI * kappa, rtol=1e-9)
    np.testing.assert_allclose(r["values"]["D"], 0, atol=1e-9)
    np.testing.assert_allclose(r["values"]["R"], 0, atol=1e-9)
    assert r["values"]["M"][0] > 0  # hot top, restrained: sagging
    assert r["reactions"][0]["moment"] == pytest.approx(EI * kappa, rel=1e-9)
    assert r["meta"]["equilibrium_error"] < 1e-9


def test_nonprismatic_fixed_deflection_uses_refined_integration():
    # A "tapered" zone between identical sections is prismatic in disguise.
    L, w = 20.0, 10.0
    girder = Section()
    model = Model(
        spans=[
            Span(length=L, zones=[Zone(end=0.5, section=0), Zone(end=1, section=0)])
        ],
        supports=["fixed", "fixed"],
        sections=[girder],
        nonprismatic=True,
        dead=[DeadLoad(w=w)],
        load_mode="dead",
    )
    ei = span_ei(model, 0)
    assert isinstance(ei, cba.SectionEI)
    r = analyse(model)
    ei_value = ei(1.0)
    assert r["max"]["D"][station(r, L / 2)] == pytest.approx(
        1000 * w * L**4 / (384 * ei_value), rel=1e-4
    )


def test_excel_reports_moment_reactions_for_integral_abutments():
    r = analyse(beam([20], ["fixed", "roller"]))
    wb = load_workbook(BytesIO(excel_bytes(r, "en")))
    sheet = wb["Reactions"]
    header = [c.value for c in sheet[1]]
    assert "Mr (kN·m, CCW +)" in header and "Type" in header  # dead only
    assert sheet.cell(2, header.index("Type") + 1).value == "fixed"
    thermal = analyse_thermal(
        Model(
            spans=[Span(length=10)],
            supports=["fixed", "fixed"],
            load_mode="thermal",
        )
    )
    wb = load_workbook(BytesIO(excel_bytes(thermal, "fr")))
    assert "Mr (kN·m, anti-horaire +)" in [c.value for c in wb["Réactions"][1]]


def test_project_round_trip_with_fixed_supports_and_plates():
    model = Model(
        supports=["fixed", "pin", "fixed"],
        nonprismatic=True,
        sections=[Section(), Section(name="S2", depth=1800, top_width=600)],
    )
    model.spans[0].zones = [
        Zone(end=0.8, section=0),
        Zone(end=1, section=0, end_section=1, profile="parabolic", plates="deep"),
    ]
    model.spans[1].zones = [
        Zone(end=0.2, section=1, end_section=0, profile="parabolic", plates="deep"),
        Zone(end=1, section=0),
    ]
    project = create_project(model, "Integral")
    assert project["schema_version"] == SCHEMA_VERSION
    reopened = validate_project(json.dumps(project))
    assert reopened["model"] == model.model_dump(mode="json")
    # v0.4 (schema 2) files still open: no plates key, default "start".
    old = json.loads(json.dumps(project))
    old["schema_version"] = 2
    for span in old["model"]["spans"]:
        for zone in span["zones"]:
            zone.pop("plates")
    old["model"]["supports"] = ["roller", "pin", "roller"]
    assert (
        validate_project(json.dumps(old))["model"]["spans"][0]["zones"][1]["plates"]
        == "start"
    )


# ------------------------------------------------------- plates (S9)


def taper_model(plates_left, plates_right):
    s1 = Section(name="S1", inertia_modifier=3.92)
    s2 = Section(
        name="S2", depth=1800, top_width=600, top_thickness=50, inertia_modifier=1.12
    )
    m = Model(nonprismatic=True, sections=[s1, s2], load_mode="live")
    m.spans[0].zones = [
        Zone(end=0.8, section=0),
        Zone(end=1, section=0, end_section=1, profile="parabolic", plates=plates_left),
    ]
    m.spans[1].zones = [
        Zone(
            end=0.2, section=1, end_section=0, profile="parabolic", plates=plates_right
        ),
        Zone(end=1, section=0),
    ]
    return m


@pytest.mark.parametrize("left,right", [("deep", "deep"), ("end", "start")])
def test_plate_options_make_swapped_haunches_mirror(left, right):
    m = taper_model(left, right)
    e1, e2 = span_ei(m, 0), span_ei(m, 1)
    L = 34.8
    for x in (27.9, 30.0, 33.0, 34.79):
        assert e1(x) == pytest.approx(e2(L - x), rel=1e-9)
    r = analyse(m)
    x = np.array(r["x"])
    v = np.array(r["min"]["M"])
    assert np.abs(v - np.interp(x[-1] - x, x, v)).max() < 1e-6 * np.abs(v).max()


def test_default_plate_option_keeps_v04_behaviour():
    m = taper_model("start", "start")
    e1, e2 = span_ei(m, 0), span_ei(m, 1)
    assert e1(34.79) / e2(0.01) > 2  # the asymmetric v0.4 interpretation


def test_plate_option_validates_flange_fit():
    thick = Section(name="T", depth=400, top_thickness=150, bottom_thickness=150)
    with pytest.raises(ValueError):
        Model(
            nonprismatic=True,
            sections=[Section(depth=280, top_thickness=20, bottom_thickness=20), thick],
            spans=[
                Span(
                    zones=[
                        Zone(
                            end=1,
                            section=0,
                            end_section=1,
                            profile="linear",
                            plates="end",
                        )
                    ]
                )
            ],
            supports=["pin", "roller"],
        )


# ------------------------------------------------ influence + traverse


def job(model, key="t"):
    return json.loads(
        dispatch(
            json.dumps(
                {"action": "analyse", "data": {"model": model.model_dump(), "job": key}}
            )
        )
    )


def test_influence_line_matches_unit_load_basis_and_reciprocity():
    model = Model(load_mode="live", supports=["fixed", "pin", "roller"])
    result = job(model, "il")
    nx = len(result["x"])
    i = station(result, 13.92)
    il = json.loads(
        dispatch(json.dumps({"action": "influence", "data": {"job": "il", "index": i}}))
    )
    assert il["x_station"] == pytest.approx(13.92)
    assert il["support"] == 1 and il["fixed"] is True
    # Maxwell–Betti: deflection at i due to unit load at j equals deflection at
    # j due to unit load at i.
    basis = cached_basis(structure_key(model))
    j = station(result, 50.0)
    d_ij = basis.unit([basis.x[j]])[0][2 * nx + i]
    d_ji = basis.unit([basis.x[i]])[0][2 * nx + j]
    assert d_ij == pytest.approx(d_ji, rel=2e-3)
    governing = il["governing_max"]
    assert governing["axles"] and governing["record"] == result["case_max"][nx + i]
    # The line itself reproduces the governing envelope value (truck case).
    xs = np.array(il["x"])
    eta = np.interp([a["x"] for a in governing["axles"]], xs, il["M"])
    assert float(eta @ [a["load"] for a in governing["axles"]]) == pytest.approx(
        result["max"]["M"][i], rel=2e-3
    )


def test_influence_rejects_bad_station_and_thermal():
    model = Model(load_mode="live")
    job(model, "bad")
    with pytest.raises(ValueError):
        dispatch(
            json.dumps({"action": "influence", "data": {"job": "bad", "index": -1}})
        )
    thermal = Model(load_mode="thermal")
    job(thermal, "th")
    with pytest.raises(ValueError, match="thermal"):
        dispatch(json.dumps({"action": "influence", "data": {"job": "th", "index": 0}}))


@pytest.mark.parametrize("direction", ["forward", "reverse"])
def test_traverse_frames_match_individual_snapshots(direction):
    model = Model(load_mode="both", supports=["fixed", "pin", "roller"])
    job(model, "tr")
    data = json.loads(
        dispatch(
            json.dumps(
                {
                    "action": "traverse",
                    "data": {"job": "tr", "direction": direction, "frames": 30},
                }
            )
        )
    )
    frames = data["frames"]
    assert len(frames) == 30
    for frame in (frames[3], frames[15], frames[27]):
        snap = json.loads(
            dispatch(
                json.dumps(
                    {
                        "action": "snapshot",
                        "data": {
                            "job": "tr",
                            "index": 0,
                            "sense": "max",
                            "position": frame["position"],
                            "direction": direction,
                        },
                    }
                )
            )
        )
        np.testing.assert_allclose(frame["M"], snap["M"], atol=1e-3)
        np.testing.assert_allclose(frame["R"], snap["R"], atol=1e-3)
        np.testing.assert_allclose(frame["Mr"], snap["Mr"], atol=1e-3)


def test_traverse_lane_case_includes_full_lane_load():
    model = Model(load_mode="live", live=LiveLoad(vehicle="CL625", case="lane"))
    job(model, "lane")
    data = json.loads(
        dispatch(json.dumps({"action": "traverse", "data": {"job": "lane"}}))
    )
    frame = data["frames"][0]
    assert frame["lane"] and frame["lane"][0]["w"] == pytest.approx(9.0)
    assert frame["record"]["case"] == "lane"
