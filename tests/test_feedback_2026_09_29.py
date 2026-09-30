"""September 29 feedback: supplementary HL-93, combined exports, isolation."""
import base64
from io import BytesIO
import json

import numpy as np
from openpyxl import load_workbook
import pytest

from quickerbridge.browser import dispatch, _results
from quickerbridge.engine import analyse, snapshot
from quickerbridge.exports import excel_bytes
from quickerbridge.models import Model
from quickerbridge.projects import create_project, validate_project


def test_two_truck_scope_reconstruction_and_factors():
    m = Model.model_validate(
        {
            "spans": [{"length": 35}, {"length": 35}],
            "supports": ["roller", "pin", "roller"],
            "load_mode": "live",
            "live": {
                "vehicle": "HL93Truck",
                "case": "lane",
                "direction": "forward",
                "two_trucks": True,
                "factor": 1.4,
                "axle_factor": 0.7,
            },
        }
    )
    ordinary = m.model_copy(deep=True)
    ordinary.live.two_trucks = False
    base, r = analyse(ordinary), analyse(m)
    assert r["min"]["M"] != base["min"]["M"]
    for sense in ("min", "max"):
        for effect in ("V", "D"):
            np.testing.assert_array_equal(r[sense][effect], base[sense][effect])
        np.testing.assert_array_equal(
            np.array(r[sense]["R"])[[0, -1]], np.array(base[sense]["R"])[[0, -1]]
        )
    np.testing.assert_array_equal(r["max"]["M"], base["max"]["M"])
    assert r["max"]["R"][1] > base["max"]["R"][1]
    for key in ("min", "max"):
        found = 0
        for i, record in enumerate(r["case_" + key]):
            if record["case"] != "hl93_two_trucks":
                continue
            found += 1
            assert record["gap"] >= 15.24 - 1e-8
            assert record["reduction"] == 0.9 and record["factor"] == 1.33
            snap = snapshot(m, record, i, key)
            flat = np.r_[snap["V"], snap["M"], snap["D"], snap["R"], snap["Mr"]]
            expected = np.r_[
                r[key]["V"], r[key]["M"], r[key]["D"], r[key]["R"], r[key]["Mr"]
            ]
            assert flat[i] == pytest.approx(expected[i], abs=1e-6)
            assert abs(snap["equilibrium"]["force"]) < 1e-7
            assert abs(snap["equilibrium"]["moment"]) < 1e-5
            if found == 2:
                break
        if key == "min":
            assert found
    reopened = validate_project(json.dumps(create_project(m, "HL93")))
    assert reopened["model"]["live"]["two_trucks"]
    # Independent three-moment solution for two equal prismatic spans.
    # A point load P at a metres from the exterior support gives
    # M_pier = -P*a*(L²-a²)/(4L²), R_pier = P*a/L - 2*M_pier/L.
    length = 35.0
    centres = np.arange(0, 2 * length + 0.01, 0.05)
    offsets = np.array([0, 14, 28]) * 0.3048
    x = centres[:, None] - offsets + offsets[-1] / 2
    a = np.minimum(x, 2 * length - x)
    on_bridge = (a >= 0) & (a <= length)
    moments = np.where(on_bridge, a * (length**2 - a**2) / (4 * length**2), 0)
    reactions = np.where(on_bridge, a / length + 2 * moments / length, 0)
    left, right = centres <= length, centres >= length
    valid = centres[right][None, :] - centres[left][:, None] >= 15.24 + offsets[-1]
    weights = np.array([35, 145, 145])
    expected = []
    for effect, lane in (
        (moments, 9.3 * length**2 / 8),
        (reactions, 1.25 * 9.3 * length),
    ):
        truck = effect @ weights
        pairs = truck[left, None] + truck[None, right]
        best = np.max(np.where(valid, pairs, -np.inf))
        expected.append(0.9 * 1.4 * (1.33 * 0.7 * best + lane))
    pier = np.argmin(np.abs(np.array(r["x"]) - length))
    assert -r["min"]["M"][pier] == pytest.approx(expected[0], rel=5e-4)
    assert r["max"]["R"][1] == pytest.approx(expected[1], rel=5e-4)


def test_modal_workbook_values_and_no_mass():
    m = Model.model_validate(
        {
            "spans": [{"length": 20}],
            "supports": ["pin", "roller"],
            "load_mode": "dead",
            "modal": {"mass_source": "custom", "mass": 20, "modes": 3},
        }
    )
    r = analyse(m)
    wb = load_workbook(BytesIO(excel_bytes(r, "en")))
    assert wb["Modes"].max_row == 4
    assert wb["Mode shapes"].max_column == 4
    assert len(wb["Mode shapes"]._charts) == 1
    assert wb["Modes"]["E2"].number_format == "0.0%"
    assert wb["Modes"]["F4"].value == pytest.approx(
        sum(wb["Modes"].cell(i, 5).value for i in range(2, 5))
    )
    assert wb["Stations"]["N2"].value == pytest.approx(0)
    assert wb["Modes"].tables and wb["Stations"].tables
    m.dead = []
    m.modal.mass_source = "dead"
    empty = load_workbook(BytesIO(excel_bytes(analyse(m), "fr")))
    assert "indisponibles" in empty["Modes"]["A2"].value
    assert "Formes modales" not in empty.sheetnames


def test_comparison_keeps_current_job_and_export_uses_current_modal_settings():
    m = Model.model_validate(
        {"spans": [{"length": 10}], "supports": ["pin", "roller"], "load_mode": "dead"}
    )
    request = lambda action, data: json.loads(
        dispatch(json.dumps({"action": action, "data": data}))
    )
    request("analyse", {"model": m.model_dump(), "job": "feedback-current"})
    current = _results["feedback-current"]
    for length in (12, 14, 16, 18):
        other = m.model_copy(deep=True)
        other.spans[0].length = length
        result = request("compare", {"model": other.model_dump()})
        assert result["x"][-1] == length
    assert _results["feedback-current"] is current
    encoded = request(
        "excel",
        {
            "job": "feedback-current",
            "lang": "fr",
            "modal": {"mass_source": "custom", "mass": 35, "modes": 3},
        },
    )
    wb = load_workbook(BytesIO(base64.b64decode(encoded)))
    assert wb["Modes"].max_row == 4
    metadata = dict(wb["Modèle"].iter_rows(values_only=True))
    assert json.loads(metadata["Model JSON"])["modal"]["mass"] == 35
