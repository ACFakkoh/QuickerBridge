"""v1.01: FT for steel grid, aluminum and transverse wood decks on steel
girders (S6-25 5.6.7, Tables 5.8 to 5.12; Annex A5.3.4, Tables A5.3.4, A5.3.5).
"""

import math
from io import BytesIO

import pytest
from openpyxl import load_workbook

import quickerbridge.distribution as D
from quickerbridge.exports import excel_bytes
from quickerbridge.engine import analyse
from quickerbridge.models import Model

coef = D.deck_coefficients


def clamp(v, lo, hi):
    return min(max(v, lo), hi)


def bridge(deck, lengths=(20.0,), **distribution):
    data = {
        "enabled": True,
        "bridge_type": deck,
        "girders": 4,
        "spacing": 2.5,
        "overhang": 1.0,
        "carriageway": 8.5,
        "skew": 20.0,
        **distribution,
    }
    return Model.model_validate(
        {
            "spans": [{"length": L} for L in lengths],
            "supports": ["pin"] + ["roller"] * len(lengths),
            "distribution": data,
        }
    )


def rows(result, **match):
    return [r for r in result["rows"] if all(r[k] == v for k, v in match.items())]


# --- Table 5.8 / A5.3.4: steel grid --------------------------------------------


@pytest.mark.parametrize(
    "t, state, n, cls, moment, shear",
    [
        (90, "ULS", 1, "AB", 2.70, 2.30),
        (90, "ULS", 3, "AB", 2.70, 2.30),
        (90, "FLS", 2, "AB", 2.70, 2.40),
        (100, "ULS", 1, "AB", 3.60, 3.0),
        (120, "ULS", 2, "AB", 3.40, 3.0),
        (120, "FLS", 2, "AB", 3.60, 3.0),
        (90, "ULS", 1, "CD", 2.70, 2.30),
        (90, "ULS", 2, "CD", 2.80, 2.50),
        (90, "FLS", 3, "CD", 2.70, 2.40),
        (120, "ULS", 1, "CD", 3.60, 3.0),
        (120, "ULS", 3, "CD", 3.60, 3.15),
        (120, "FLS", 2, "CD", 3.60, 3.0),
    ],
)
def test_steel_grid_table(t, state, n, cls, moment, shear):
    for girder in ("interior", "exterior"):
        assert coef("steel_grid", state, girder, "moment", n, 20, 2.5, t, cls) == (
            pytest.approx(moment),
            0.0,
            1.0,
        )
    assert coef("steel_grid", state, "interior", "shear", n, 20, 2.5, t, cls) == (
        pytest.approx(shear),
        0.0,
        1.0,
    )


# --- Tables 5.9 to 5.11: aluminum ------------------------------------------------


def test_aluminum_uls_moment():
    dt, lam, gc = coef("aluminum", "ULS", "interior", "moment", 2, 20, 2.0, 0, "AB")
    assert dt == pytest.approx(3.8 - 4.4 / 5)
    assert lam == pytest.approx(0.10 - 0.25 / 20)
    assert gc == 1.0
    # Lower bound 2.55 (short Le).
    assert coef("aluminum", "ULS", "interior", "moment", 2, 3, 2.0, 0, "AB")[
        0
    ] == pytest.approx(max(3.8 - 4.4 / math.sqrt(8), 2.55))
    # Classes C and D: λ = 0 (A5.3.4.3).
    assert coef("aluminum", "ULS", "interior", "moment", 2, 20, 2.0, 0, "CD")[1] == 0


def test_aluminum_shear():
    dt, lam, gc = coef("aluminum", "ULS", "interior", "shear", 2, 16, 1.5, 0, "AB")
    assert dt == pytest.approx(min(3.9 - 3.6 / 4, 3.1))
    assert lam == 0
    assert gc == pytest.approx(0.75**0.25)
    assert coef("aluminum", "ULS", "interior", "shear", 2, 60, 2.0, 0, "AB")[
        0
    ] == pytest.approx(3.1)
    dt, _, gc = coef("aluminum", "FLS", "interior", "shear", 2, 16, 2.5, 0, "AB")
    assert dt == pytest.approx(5.2 - 5.4 / 4)
    assert gc == 1.0


def test_aluminum_fls_moment_gamma_c():
    dt, lam, gc = coef("aluminum", "FLS", "interior", "moment", 2, 20, 2.0, 0, "AB")
    assert dt == pytest.approx(4 - 2.5 / math.sqrt(20))
    assert lam == pytest.approx(0.05)
    assert gc == pytest.approx(1 + 0.9 * (0.3 * 2 - 0.36) * 10 / 40)
    # One lane, short Le, small S: γc = 1.
    assert coef("aluminum", "FLS", "interior", "moment", 1, 20, 2.0, 0, "AB")[2] == 1
    assert coef("aluminum", "FLS", "interior", "moment", 2, 8, 2.0, 0, "AB")[2] == 1
    assert coef("aluminum", "FLS", "interior", "moment", 2, 20, 1.2, 0, "AB")[2] == 1
    # Exterior girder: Table 5.10 covers interior girders only.
    assert coef("aluminum", "FLS", "exterior", "moment", 2, 20, 2.0, 0, "AB")[2] == 1
    # Beyond Table 5.10 (Le > 30 m, S > 3 m): capped values.
    assert coef("aluminum", "FLS", "interior", "moment", 2, 50, 4.0, 0, "AB")[
        2
    ] == pytest.approx(1 + 0.9 * (0.3 * 3 - 0.36) * 20 / 40)


# --- Table 5.12 / A5.3.5: transverse wood decks ---------------------------------


def test_wood_laminated_140_ab():
    c = lambda st, n, le, eff="moment": coef(  # noqa: E731
        "wood_laminated", st, "interior", eff, n, le, 2.5, 140, "AB"
    )
    assert c("ULS", 1, 10) == (pytest.approx(clamp(3.70 - 0.3, 3.0, 3.40)), 0.0, 1.0)
    assert c("ULS", 2, 30)[0] == pytest.approx(clamp(4.40 - 70 / 70, 3.0, 3.40))
    assert c("ULS", 2, 5)[0] == pytest.approx(clamp(4.40 - 70 / 45, 3.0, 3.40))
    assert c("FLS", 2, 10)[0] == pytest.approx(clamp(3.90 - 0.3, 3.40, 3.75))
    assert c("ULS", 3, 10, "shear")[0] == pytest.approx(3.0)


def test_wood_laminated_290_ab():
    c = lambda st, n, le: coef(  # noqa: E731
        "wood_laminated", st, "exterior", "moment", n, le, 2.5, 290, "AB"
    )
    assert c("ULS", 1, 10) == (pytest.approx(3.50), 0.0, 1.0)
    dt, lam, _ = c("ULS", 2, 10)
    assert dt == pytest.approx(clamp(2.90 + 0.4, 3.0, 3.50))
    assert lam == pytest.approx(0.10 - 0.025)
    dt, lam, _ = c("FLS", 2, 10)
    assert dt == pytest.approx(clamp(3.90 - 0.3, 3.40, 3.75))
    assert lam == pytest.approx(0.10 - 0.015)


def test_wood_laminated_interpolation_and_clamp():
    args = ("ULS", "interior", "moment", 2, 10, 2.5)
    a = coef("wood_laminated", *args, 140, "AB")
    b = coef("wood_laminated", *args, 290, "AB")
    m = coef("wood_laminated", *args, 215, "AB")
    assert m[0] == pytest.approx((a[0] + b[0]) / 2)
    assert m[1] == pytest.approx((a[1] + b[1]) / 2)
    assert coef("wood_laminated", *args, 100, "AB") == a
    assert coef("wood_laminated", *args, 400, "AB") == b


def test_wood_laminated_cd():
    c = lambda st, n, le, t, eff="moment": coef(  # noqa: E731
        "wood_laminated", st, "interior", eff, n, le, 2.5, t, "CD"
    )
    assert c("ULS", 2, 30, 140)[0] == pytest.approx(clamp(4.65 - 74 / 70, 3.15, 3.60))
    assert c("ULS", 3, 10, 140, "shear")[0] == pytest.approx(3.15)
    assert c("ULS", 1, 10, 140, "shear")[0] == pytest.approx(3.0)
    dt, lam, _ = c("ULS", 2, 10, 290)
    assert dt == pytest.approx(clamp(3.05 + 0.42, 3.15, 3.70))
    assert lam == pytest.approx(0.10 - 0.025)
    # n above 3 uses the "2 and 3" row.
    assert c("ULS", 4, 10, 290) == c("ULS", 3, 10, 290)


@pytest.mark.parametrize(
    "state, n, cls, moment, shear",
    [
        ("ULS", 1, "AB", 2.40, 2.40),
        ("ULS", 2, "AB", 2.55, 2.55),
        ("FLS", 3, "AB", 2.40, 2.40),
        ("ULS", 1, "CD", 2.40, 2.40),
        ("ULS", 3, "CD", 2.65, 2.65),
        ("FLS", 2, "CD", 2.40, 2.40),
    ],
)
def test_wood_plank(state, n, cls, moment, shear):
    assert coef("wood_plank", state, "interior", "moment", n, 20, 2.5, 0, cls) == (
        pytest.approx(moment),
        0.0,
        1.0,
    )
    assert coef("wood_plank", state, "interior", "shear", n, 20, 2.5, 0, cls)[
        0
    ] == pytest.approx(shear)


# --- Lever rule for the exterior-girder shear ------------------------------------


def test_lever_rule_inside_and_overhang():
    # B = 9.5, curb face 0.5 m from the deck edge, wheels at 1.1 and 2.9 m,
    # exterior girder at 1.0 m, first interior girder at 3.5 m.
    assert D.lever_exterior(4, 2.5, 1.0, 8.5, 0.6) == pytest.approx(
        0.5 * (1 - 0.1 / 2.5) + 0.5 * (1 - 1.9 / 2.5)
    )
    # Outer wheel on the overhang (outside the exterior girder).
    assert D.lever_exterior(4, 2.5, 1.0, 8.5, 0.2) == pytest.approx(
        0.5 * (1 + 0.3 / 2.5) + 0.5 * (1 - 1.5 / 2.5)
    )
    # Outer wheel 0.4 m outside the exterior girder, second wheel beyond the
    # first interior girder (no contribution).
    assert D.lever_exterior(4, 1.2, 1.0, 5.6, 0.6) == pytest.approx(
        0.5 * (1 + 0.4 / 1.2)
    )


# --- Whole calculation -------------------------------------------------------------


@pytest.mark.parametrize(
    "deck", ["steel_grid", "aluminum", "wood_laminated", "wood_plank"]
)
def test_deck_fraction(deck):
    r = D.truck_fraction(bridge(deck))
    assert r["kind"] == deck
    assert all(row["Fs"] == 1.0 and row["gamma_e"] == 0.0 for row in r["rows"])
    assert all(z["Fs"] == 1.0 for z in r["zones"])
    assert r["derived"]["table"]
    for state in ("ULS", "FLS"):
        (inner,) = rows(r, state=state, girder="interior", effect="shear")
        (outer,) = rows(r, state=state, girder="exterior", effect="shear")
        lever = r["derived"]["lever"]
        assert outer["lever"] == pytest.approx(lever)
        assert outer["FT"] == pytest.approx(max(lever, inner["FT"]))
        assert outer["method"] == ("lever" if lever > inner["FT"] else "interior")
        for row in rows(r, state=state):
            calc = 2.5 / (
                row["DT"] * row["gamma_c"] * (1 + r["derived"]["mu"] * row["lambda"])
            )
            if not (row["girder"] == "exterior" and row["effect"] == "shear"):
                assert row["FT"] == pytest.approx(max(calc, row["FT_min"]))


def test_wood_thickness_warning():
    assert (
        "wood_thickness"
        in D.truck_fraction(bridge("wood_laminated", deck_thickness=100))["warnings"]
    )
    assert (
        "wood_thickness"
        not in D.truck_fraction(bridge("wood_laminated", deck_thickness=200))[
            "warnings"
        ]
    )


def test_aluminum_range_warning():
    r = D.truck_fraction(bridge("aluminum", lengths=(60.0,), spacing=3.5))
    assert "alu_range" in r["warnings"]


def test_v10_project_defaults():
    m = Model.model_validate(
        {
            "spans": [{"length": 20}],
            "supports": ["pin", "roller"],
            "distribution": {"enabled": True},
        }
    )
    assert m.distribution.deck_thickness == 150
    assert m.distribution.wheel_edge == 0.6
    assert D.truck_fraction(m)["kind"] == "slab_on_girder"


def test_deck_analysis_and_excel():
    m = bridge("wood_laminated", deck_thickness=215, apply=True, girder="exterior")
    result = analyse(m)
    assert result["ft"]["zones"]
    assert all(z["Fs"] == 1.0 for z in result["ft"]["zones"])
    book = load_workbook(BytesIO(excel_bytes(result, "en", model=m)))
    text = " ".join(
        str(c.value)
        for ws in book.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    )
    assert "215" in text and "0.6" in text.replace(",", ".")
