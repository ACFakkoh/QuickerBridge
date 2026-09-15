from io import BytesIO, StringIO
import csv

import numpy as np
import pytest
from openpyxl import load_workbook
import pycba as cba

from quickerbridge.models import Model, Span, Section, LiveLoad, DeadLoad, Zone
from quickerbridge.sections import properties, span_ei
from quickerbridge.loads import (
    axle_groups,
    dynamic_factor,
    lane_parameters,
    vehicle_data,
    vehicle_variants,
)
from quickerbridge.engine import analyse, snapshot, Basis
from quickerbridge.exports import csv_bytes, excel_bytes


def simple(length=10, **kwargs):
    return Model(spans=[Span(length=length)], supports=["pin", "roller"], **kwargs)


def test_section_parallel_axis_rectangle():
    section = Section(
        E=30,
        depth=1000,
        top_width=300,
        bottom_width=300,
        web_thickness=300,
        top_thickness=50,
        bottom_thickness=50,
    )
    p = properties(section)
    assert p["A"] == pytest.approx(0.3)
    assert p["centroid"] == pytest.approx(500)
    assert p["I"] == pytest.approx(0.3 / 12)
    assert p["EI"] == pytest.approx(30e6 * 0.3 / 12)


def test_s6_dynamic_groups_exact_original_numbering():
    assert dynamic_factor([4]) == 1.4
    assert dynamic_factor([1, 2]) == 1.3
    assert dynamic_factor([1, 2, 3]) == 1.3
    assert dynamic_factor([3, 4, 5]) == 1.25
    assert dynamic_factor([1, 2, 3, 4, 5]) == 1.25
    assert dynamic_factor([1, 2, 3], False) == 1
    for name in ("CL625", "CL750QC"):
        assert len(axle_groups(LiveLoad(vehicle=name))) == 31


def test_pycba_us_vehicle_definitions_and_factors():
    truck = LiveLoad(vehicle="HL93Truck")
    tandem = LiveLoad(vehicle="HL93Tandem")
    cooper = LiveLoad(vehicle="Cooper", cooper_e=80)
    np.testing.assert_allclose(vehicle_data(truck)[0], [35, 145, 145])
    np.testing.assert_allclose(vehicle_data(tandem)[0], [110, 110])
    np.testing.assert_allclose(
        vehicle_data(cooper)[0],
        cba.VehicleLibrary.US.get_cooper(80).axw,
    )
    variants = vehicle_variants(truck)
    assert len(variants) == 20
    assert variants[0][2]["rear_spacing"] == pytest.approx(4.3)
    assert variants[-1][2]["rear_spacing"] == pytest.approx(9.0)
    assert dynamic_factor([1, 2, 3], vehicle="HL93Truck") == pytest.approx(1.33)
    assert dynamic_factor([1, 2], vehicle="Cooper") == 1
    assert (
        len(axle_groups(truck))
        == len(axle_groups(tandem))
        == len(axle_groups(cooper))
        == 1
    )
    assert lane_parameters(truck) == (9.3, 1.0, "full")
    assert lane_parameters(cooper)[0] == pytest.approx(116.75122309711286)


def test_hl93_full_lane_applies_im_to_axles_only_and_cooper_snapshot_balances():
    base = dict(
        spans=[Span(length=12)],
        supports=["pin", "roller"],
        sections=[Section(kind="ei", EI=20_000_000)],
        load_mode="live",
    )
    dynamic = Model(**(base | {"live": LiveLoad(vehicle="HL93Tandem", case="lane")}))
    nominal = dynamic.model_copy(deep=True)
    nominal.live.dynamic = False
    a, b = analyse(dynamic), analyse(nominal)
    assert max(a["max"]["M"]) > max(b["max"]["M"])
    assert any(
        case["case"] == "lane" and case["factor"] == pytest.approx(1.33)
        for case in a["case_max"]
    )

    cooper = Model(**(base | {"live": LiveLoad(vehicle="Cooper", case="lane")}))
    envelope = analyse(cooper)
    extreme = envelope["extrema"][0]
    view = snapshot(cooper, extreme, extreme["index"], extreme["sense"])
    index = extreme["index"] % len(envelope["x"])
    assert view[extreme["response"]][index] == pytest.approx(extreme["value"], abs=1e-7)
    assert abs(view["equilibrium"]["force"]) < 1e-6
    assert abs(view["equilibrium"]["moment"]) < 1e-4


def test_single_span_udl_exact_and_deflection():
    model = simple(load_mode="dead")
    r = analyse(model)
    ei = properties(model.sections[0])["EI"]
    assert r["reactions"][0]["max"] == pytest.approx(50)
    assert max(r["max"]["M"]) == pytest.approx(10 * 10**2 / 8)
    assert max(r["max"]["D"]) == pytest.approx(
        5 * 10 * 10**4 / (384 * ei) * 1000, rel=3e-5
    )
    assert r["max"]["D"][0] == 0
    assert r["max"]["D"][-1] == 0


def test_two_span_continuity_and_both_shear_sides():
    model = Model(spans=[Span(length=30), Span(length=30)], load_mode="dead")
    r = analyse(model)
    np.testing.assert_allclose([v["max"] for v in r["reactions"]], [112.5, 375, 112.5])
    assert min(r["max"]["M"]) == pytest.approx(-1125)
    at_pier = [i for i, x in enumerate(r["x"]) if x == 30]
    assert len(at_pier) == 2
    assert r["max"]["V"][at_pier[1]] - r["max"]["V"][at_pier[0]] == pytest.approx(375)
    np.testing.assert_allclose(r["max"]["D"], r["max"]["D"][::-1], atol=1e-10)
    assert all(r["max"]["D"][i] == 0 for i in at_pier)


def test_point_snapshot_matches_closed_form_and_equilibrium():
    live = LiveLoad(
        vehicle="custom", weights=[100], spacings=[], dynamic=False, case="truck"
    )
    model = simple(load_mode="live", live=live)
    r = snapshot(
        model,
        {
            "case": "truck",
            "axles": [1],
            "factor": 1,
            "direction": "forward",
            "position": 5,
        },
    )
    ei = properties(model.sections[0])["EI"]
    assert max(r["M"]) == pytest.approx(250)
    assert max(r["D"]) == pytest.approx(100 * 10**3 / (48 * ei) * 1000, rel=3e-5)
    assert abs(r["equilibrium"]["force"]) < 1e-9
    assert abs(r["equilibrium"]["moment"]) < 1e-9
    graph = r["plot"]
    at_axle = [i for i, x in enumerate(graph["x"]) if x == 5]
    assert len(at_axle) == 2
    assert graph["V"][at_axle[1]] - graph["V"][at_axle[0]] == pytest.approx(-100)


def test_short_span_shear_limit_and_nonround_midspan():
    live = LiveLoad(
        vehicle="custom", weights=[100], spacings=[], dynamic=True, case="truck"
    )
    model = simple(length=1.3, load_mode="live", live=live)
    r = analyse(model)
    assert max(r["max"]["V"]) == pytest.approx(140, rel=1e-6)
    assert max(r["max"]["M"]) == pytest.approx(140 * 1.3 / 4, rel=1e-6)


def test_dead_load_superposition_and_partial_udl():
    a = simple(load_mode="dead", dead=[DeadLoad(w=3, start=0.1, end=0.6)])
    b = simple(load_mode="dead", dead=[DeadLoad(w=7, start=0.1, end=0.6)])
    ab = simple(load_mode="dead", dead=a.dead + b.dead)
    ra, rb, rab = analyse(a), analyse(b), analyse(ab)
    for key in ("V", "M", "D", "R"):
        np.testing.assert_allclose(
            np.array(ra["max"][key]) + rb["max"][key], rab["max"][key], atol=1e-10
        )
    assert sum(rab["max"]["R"]) == pytest.approx(50)


def test_lane_has_no_dynamic_amplification_and_fraction_matters():
    a = simple(
        length=12,
        load_mode="live",
        live=LiveLoad(vehicle="CL750QC", case="lane", lane_fraction=0.63, dynamic=True),
    )
    b = a.model_copy(deep=True)
    b.live.dynamic = False
    ra, rb = analyse(a), analyse(b)
    np.testing.assert_allclose(ra["max"]["M"], rb["max"]["M"])
    b.live.lane_fraction = 0.8
    rc = analyse(b)
    assert max(rc["max"]["M"]) > max(ra["max"]["M"])
    assert all(c["factor"] == 1 for c in rc["case_max"])


def test_user_load_factors_and_maintenance_vehicle():
    dead = simple(load_mode="dead", dead=[DeadLoad(w=10, factor=2)])
    assert max(analyse(dead)["max"]["M"]) == pytest.approx(250)

    base = simple(
        length=10.13,
        load_mode="live",
        live=LiveLoad(vehicle="Maintenance", case="truck", direction="both"),
    )
    factored = base.model_copy(deep=True)
    factored.live.factor = 2
    factored.live.axle_factor = 3
    a, b = analyse(base), analyse(factored)
    for sense in ("min", "max"):
        for effect in ("V", "M", "D", "R"):
            np.testing.assert_allclose(b[sense][effect], 6 * np.array(a[sense][effect]))
    np.testing.assert_allclose(vehicle_data(base.live)[0], [24, 56])
    np.testing.assert_allclose(vehicle_data(base.live)[1], [0, 2])
    assert lane_parameters(base.live) == (0, 1, "none")
    assert axle_groups(base.live)[0]["factor"] == 1


def test_canadian_lane_is_full_deck_and_standard_grid_is_symmetric():
    live = LiveLoad(vehicle="CL750QC", case="lane", lane_fraction=0.8, factor=1.5)
    assert lane_parameters(live) == (12.6, 0.8, "full")
    model = simple(length=10.13, load_mode="live", live=live)
    result = analyse(model)
    view = snapshot(model, result["extrema"][0], result["extrema"][0]["index"], "max")
    assert view["lane"] == [{"start": 0.0, "end": 10.13, "w": 18.9}]
    assert all(
        case["case"] != "lane" or not case["axles"] or case["axles"] == [1, 2, 3, 4, 5]
        for case in result["case_min"] + result["case_max"]
    )

    maintenance = simple(
        length=10.13,
        load_mode="live",
        live=LiveLoad(vehicle="Maintenance", case="truck", direction="both"),
    )
    envelope = analyse(maintenance)
    for effect in ("M", "D"):
        np.testing.assert_allclose(
            envelope["max"][effect], envelope["max"][effect][::-1], atol=1e-10
        )


def test_governing_snapshot_reproduces_each_extreme():
    model = Model()
    r = analyse(model)
    for e in r["extrema"]:
        s = snapshot(model, e, e["index"], e["sense"])
        i = e["index"] % len(r["x"])
        assert s[e["response"]][i] == pytest.approx(e["value"], abs=1e-7)
        assert abs(s["equilibrium"]["force"]) < 1e-6
        assert abs(s["equilibrium"]["moment"]) < 1e-4


def test_constant_zone_equals_prismatic():
    a = simple(load_mode="dead")
    b = a.model_copy(deep=True)
    b.nonprismatic = True
    b.spans[0].zones = [Zone(end=0.35, section=0), Zone(end=1, section=0)]
    ra, rb = analyse(a), analyse(b)
    for key in ("V", "M", "D", "R"):
        np.testing.assert_allclose(ra["max"][key], rb["max"][key], rtol=1e-8, atol=1e-8)


def test_taper_and_step_against_refined_pycba_members():
    # Independent mesh representation: many constant-EI members with free
    # intermediate nodes, compared to one SectionEI member. This catches
    # erroneous span/zone coordinates and false supports at section changes.
    a, b = Section(), Section(depth=2600)
    model = simple(load_mode="dead", nonprismatic=True, sections=[a, b])
    model.spans[0].zones = [
        Zone(end=0.35, section=1),
        Zone(end=1, section=1, end_section=0, profile="linear"),
    ]
    r = analyse(model)
    sec = span_ei(model, 0)
    n = 200
    dx = 10 / n
    ba = cba.BeamAnalysis(
        [dx] * n,
        sec((np.arange(n) + 0.5) * dx),
        supports=["pin"] + ["free"] * (n - 1) + ["roller"],
    )
    ba.set_loads([[i + 1, 1, 10] for i in range(n)])
    ba.analyze(npts=4)
    # The midpoint is a node, so compare solved nodal displacement directly.
    expected = -ba.beam_results.D[2 * (n // 2)] * 1000
    i = np.argmin(abs(np.array(r["x"]) - 5))
    assert r["max"]["D"][i] == pytest.approx(expected, rel=0.004)


def test_custom_seven_and_five_spans_dead_shape():
    assert (
        len(
            axle_groups(
                LiveLoad(vehicle="custom", weights=[80] * 7, spacings=[1.2] * 6)
            )
        )
        == 127
    )
    model = Model(
        spans=[Span(length=10)] * 5, supports=["pin"] + ["roller"] * 5, load_mode="dead"
    )
    r = analyse(model)
    assert len(r["reactions"]) == 6
    assert len(r["table"]) == 55
    assert sum(r["max"]["R"]) == pytest.approx(500)


def test_exports_consistent_reactions_once_and_french():
    r = analyse(Model(spans=[Span(length=30), Span(length=30)], load_mode="dead"))
    data = list(csv.reader(StringIO(csv_bytes(r, "fr").decode("utf-8-sig"))))
    assert data[0][0] == "Travée"
    assert len(data) == 23
    assert sum(bool(row[-1]) for row in data[1:]) == 3
    wb = load_workbook(BytesIO(excel_bytes(r, "fr")))
    assert "Réactions" in wb.sheetnames
    assert wb["Réactions"]["D3"].value == pytest.approx(375)
    assert wb["Stations"].max_row == 23


def test_validation_rejects_invalid_geometry_and_zones():
    with pytest.raises(ValueError):
        Section(depth=20)
    with pytest.raises(ValueError):
        Model(
            nonprismatic=True,
            spans=[Span(zones=[Zone(end=0.8, section=0)])],
            supports=["pin", "roller"],
        )
    with pytest.raises(ValueError):
        LiveLoad(weights=[100, 100], spacings=[])
