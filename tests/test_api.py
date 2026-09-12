import time
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from quickerbridge.server import app


def test_local_api_jobs_snapshot_exports_and_validation():
    with TestClient(app) as client:
        assert client.get("/").status_code == 200
        assert client.get("/api/health").json()["app"] == "QuickerBridge"
        assert (
            client.post(
                "/api/jobs", json={}, headers={"Origin": "https://example.com"}
            ).status_code
            == 403
        )
        model = client.get("/api/defaults").json()
        invalid = {**model, "supports": ["pin"]}
        assert client.post("/api/jobs", json=invalid).status_code == 422
        model["load_mode"] = "dead"
        response = client.post("/api/jobs", json=model)
        assert response.status_code == 200
        key = response.json()["id"]
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            result = client.get("/api/jobs/" + key).json()
            if result["status"] in ("complete", "error"):
                break
            time.sleep(0.01)
        assert result["status"] == "complete"
        snap = client.post("/api/snapshot", json={"job": key, "index": 0}).json()
        assert len(snap["R"]) == 3
        assert len(snap["plot"]["x"]) >= len(snap["x"])
        assert client.get(f"/api/export/{key}/csv?lang=fr").content.startswith(
            b"\xef\xbb\xbf"
        )
        book = client.get(f"/api/export/{key}/xlsx?lang=fr")
        assert book.status_code == 200
        assert "Réactions" in load_workbook(BytesIO(book.content)).sheetnames
        assert client.get(f"/api/export/{key}/xlsx?lang=xx").status_code == 422
