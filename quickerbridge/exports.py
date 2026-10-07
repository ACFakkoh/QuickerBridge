"""CSV and Excel exports from the same immutable result shown in the viewer."""

import csv
from io import BytesIO, StringIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import ScatterChart, Reference, Series

from .version import APP_VERSION, AUTHOR, RELEASE_DATE


WARNING = {
    "en": "Indicative preliminary values only — not a substitute for detailed design.",
    "fr": "Valeurs préliminaires indicatives seulement — ne remplacent pas une conception détaillée.",
}
HEADERS = {
    "en": [
        "Span",
        "Station",
        "x (m)",
        "Local x (m)",
        "Side",
        "V min (kN)",
        "V max (kN)",
        "M min (kN·m)",
        "M max (kN·m)",
        "Deflection min (mm, down +)",
        "Deflection max (mm, down +)",
        "R min (kN, up +)",
        "R max (kN, up +)",
    ],
    "fr": [
        "Travée",
        "Station",
        "x (m)",
        "x local (m)",
        "Côté",
        "V min (kN)",
        "V max (kN)",
        "M min (kN·m)",
        "M max (kN·m)",
        "Flèche min (mm, bas +)",
        "Flèche max (mm, bas +)",
        "R min (kN, haut +)",
        "R max (kN, haut +)",
    ],
}
THERMAL_HEADERS = {
    "en": [
        "Span",
        "Station",
        "x (m)",
        "Local x (m)",
        "Side",
        "V (kN)",
        "M (kN·m)",
        "Deflection (mm, down +)",
        "R (kN, up +)",
    ],
    "fr": [
        "Travée",
        "Station",
        "x (m)",
        "x local (m)",
        "Côté",
        "V (kN)",
        "M (kN·m)",
        "Flèche (mm, bas +)",
        "R (kN, haut +)",
    ],
}


def rows(result, language):
    seen = set()
    for row in result["table"]:
        values = list(row.values())
        if language == "fr":
            values[4] = {"left": "gauche", "right": "droite"}[values[4]]
        reaction = next(
            (r for r in result["reactions"] if abs(r["x"] - row["x"]) < 1e-8), None
        )
        if reaction and reaction["support"] not in seen:
            values.extend([reaction["min"], reaction["max"]])
            seen.add(reaction["support"])
        else:
            values.extend([None, None])
        yield values


def single_case(result):
    """One value per effect: imposed deformation or permanent loads alone."""
    return (
        result.get("kind") == "thermal"
        or result.get("model", {}).get("load_mode") == "dead"
    )


def thermal_rows(result, language):
    """Rows of a single-case result (no min/max envelope)."""
    seen = set()
    for row in result["table"]:
        values = [
            row[key] if key in row else row[key + "_max"]
            for key in ("span", "station", "x", "local_x", "side", "V", "M", "D")
        ]
        if language == "fr":
            values[4] = {"left": "gauche", "right": "droite"}[values[4]]
        reaction = next(
            (r for r in result["reactions"] if abs(r["x"] - row["x"]) < 1e-8), None
        )
        if reaction and reaction["support"] not in seen:
            values.append(reaction.get("value", reaction.get("max")))
            seen.add(reaction["support"])
        else:
            values.append(None)
        yield values


def csv_bytes(result, language="en"):
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    thermal = single_case(result)
    writer.writerow((THERMAL_HEADERS if thermal else HEADERS)[language])
    writer.writerows((thermal_rows if thermal else rows)(result, language))
    return stream.getvalue().encode("utf-8-sig")


def distribution_sheet(wb, model, language):
    """S6-25 truck load fraction FT (slab-on-girder) as one flat table."""
    from .distribution import truck_fraction

    fr = language == "fr"
    data = truck_fraction(model)
    sheet = wb.create_sheet("Facteur d'essieu FT" if fr else "Truck fraction FT")
    sheet.append(
        [
            "État limite" if fr else "Limit state",
            "Poutre" if fr else "Girder",
            "Effet" if fr else "Effect",
            "Signe" if fr else "Sign",
            "Lieu" if fr else "Location",
            "Le (m)",
            "Règle Le" if fr else "Le rule",
            "DT (m)",
            "λ",
            "γc",
            "γe",
            "FT calculé" if fr else "FT computed",
            "FT min",
            "FT",
            "Fs",
            "FT × Fs",
        ]
    )
    names = {
        "ULS": "ÉLUL / ÉLUT1" if fr else "ULS / SLS1",
        "FLS": "ÉLF / ÉLUT2" if fr else "FLS / SLS2",
        "interior": "intérieure" if fr else "interior",
        "exterior": "extérieure" if fr else "exterior",
        "moment": "moment",
        "shear": "cisaillement" if fr else "shear",
    }
    for r in data["rows"]:
        kind, number = r["where"].split(":")
        where = (
            f"{'Travée' if fr else 'Span'} {number}"
            if kind == "span"
            else f"{'Appui' if fr else 'Support'} {number}"
        )
        sheet.append(
            [
                names[r["state"]],
                names[r["girder"]],
                names[r["effect"]],
                ("positif" if fr else "positive")
                if r["sign"] == "+"
                else ("négatif" if fr else "negative"),
                where,
                r["Le"],
                r["Le_rule"],
                r["DT"],
                r["lambda"],
                r["gamma_c"],
                r["gamma_e"],
                r["FT_calc"],
                r["FT_min"],
                r["FT"],
                r["Fs"],
                r["FT_Fs"],
            ]
        )
    d = data["derived"]
    inputs = data["inputs"]
    sheet.append([])
    if "Be" in d:  # slab or voided slab: effects per metre of width
        items = [
            ("Type", inputs["bridge_type"]),
            ("B (m)", d["B"]),
            ("Be (m)", d["Be"]),
        ]
        if inputs["bridge_type"] == "voided_slab":
            items.append(("S (m)", inputs["spacing"]))
        items.append(
            ("Unité / Unit", "par mètre de largeur" if fr else "per metre of width")
        )
    else:
        items = [
            ("N", inputs["girders"]),
            ("S (m)", inputs["spacing"]),
            ("Sc (m)", inputs["overhang"]),
            ("B (m)", d["B"]),
            ("DVE (m)", d["DVE"]),
        ]
    items += [
        ("Wc (m)", inputs["carriageway"]),
        ("ψ (°)", inputs["skew"]),
        ("n", d["n"]),
        ("RL", d["RL"]),
        ("We (m)", d["We"]),
        ("μ", d["mu"]),
    ]
    # v0.9.6: every parameter of the simplified method used.
    items += [
        ("Classe de route" if fr else "Road class", inputs["road_class"]),
        ("Poutre choisie" if fr else "Selected girder", names[inputs["girder"]]),
        (
            "État limite choisi" if fr else "Selected limit state",
            names[inputs["state"]],
        ),
        (
            "FT appliqué aux enveloppes" if fr else "FT applied to the envelopes",
            ("oui" if fr else "yes") if inputs["apply"] else ("non" if fr else "no"),
        ),
        ("Bordure / Curb (m)", d.get("curb")),
        ("FT min ÉLUL" if fr else "FT min ULS", d["minimum"]["ULS"]),
        ("FT min ÉLF" if fr else "FT min FLS", d["minimum"]["FLS"]),
        ("h gauche (m)" if fr else "h left (m)", inputs["h_left"]),
        ("h droite (m)" if fr else "h right (m)", inputs["h_right"]),
    ]
    for s in d["spans"]:
        items.append((f"{'Travée' if fr else 'Span'} {s['span']} · L (m)", s["L"]))
        items.append((f"{'Travée' if fr else 'Span'} {s['span']} · Fs", s["Fs"]))
    for z in d["positive"]:
        items.append(
            (
                f"Le M+ {'travée' if fr else 'span'} {z['span']} (m) · {z['rule']}",
                z["Le"],
            )
        )
    for z in d["negative"]:
        items.append(
            (
                f"Le M− {'appui' if fr else 'support'} {z['support']} (m) · {z['rule']}",
                z["Le"],
            )
        )
    for zone in data.get("zones", []):
        if isinstance(zone, dict) and {"FT_M", "FT_V"} <= set(zone):
            where = zone.get("where", "")
            items.append(
                (
                    f"Zone {where} · FT M / FT V",
                    f"{zone['FT_M']:.4f} / {zone['FT_V']:.4f}",
                )
            )
    for label, value in items:
        sheet.append([label, value])


def excel_bytes(result, language="en", model=None, info=None):
    from .models import Model
    from .modal import analyse_modal

    model = model or Model.model_validate(result["model"])
    thermal = result.get("kind") == "thermal"
    single = single_case(result)
    wb = Workbook()
    ws = wb.active
    ws.title = "Stations"
    ws.append((THERMAL_HEADERS if single else HEADERS)[language])
    for row in (thermal_rows if single else rows)(result, language):
        ws.append(row)
    rx = wb.create_sheet("Réactions" if language == "fr" else "Reactions")
    rx.append(
        [
            "Appui" if language == "fr" else "Support",
            "x (m)",
            "R (kN)" if single else "R min (kN)",
        ]
    )
    if not single:
        rx.cell(1, 4, "R max (kN)")
    fixed = any(r.get("type") in ("fixed", "spring") for r in result["reactions"])
    if fixed:
        # Integral (fixed) abutments also carry a moment reaction, CCW +.
        column = 4 if single else 5
        labels = (
            ["Me (kN·m, anti-horaire +)"] if language == "fr" else ["Me (kN·m, CCW +)"]
        )
        if not single:
            labels = [
                label.replace("Me", name)
                for name in ("Me min", "Me max")
                for label in labels
            ]
        for offset, label in enumerate(labels):
            rx.cell(1, column + offset, label)
        rx.cell(1, column + len(labels), "Type")
    for r in result["reactions"]:
        row = (
            [r["support"], r["x"], r.get("value", r.get("max"))]
            if single
            else [r["support"], r["x"], r["min"], r["max"]]
        )
        if fixed:
            row += (
                [r.get("moment", r.get("moment_max", 0.0))]
                if single
                else [r.get("moment_min", 0.0), r.get("moment_max", 0.0)]
            )
            kind = r.get("type", "")
            if kind == "spring":
                kind = f"spring k={r['k']:,.0f} kN·m/rad ({100 * r['fixity']:.0f}%)"
            row.append(kind)
        rx.append(row)
        split = r.get("split")
        if split:
            # v0.9.8 split pier: one bearing line each side of the deck joint,
            # R left = −V(x⁻), R right = V(x⁺).
            lo, hi = result["min"]["V"], result["max"]["V"]
            i, j = split["left"], split["right"]
            sides = ("gauche", "droite") if language == "fr" else ("left", "right")
            for name, values in (
                (sides[0], [-hi[i], -lo[i]]),
                (sides[1], [lo[j], hi[j]]),
            ):
                extra = [values[1]] if single else values
                if fixed:
                    extra += [""] * (1 if single else 2) + ["split"]
                rx.append([f"{r['support']} {name}", r["x"], *extra])
    if not single:
        cases = wb.create_sheet(
            "Cas déterminants" if language == "fr" else "Governing cases"
        )
        cases.append(
            [
                "x (m)",
                "Side / Côté",
                "Effect / Effet",
                "Min/Max",
                "Value / Valeur",
                "Case / Cas",
                "Axles / Essieux",
                "DLA factor / Facteur CMD",
                "Load factor / Facteur de charge",
                "Axle factor / Facteur d’essieu",
                "Front axle x / x essieu avant (m)",
                "Direction",
            ]
        )
        nx = len(result["x"])
        for effect, offset in (("V", 0), ("M", nx), ("D", 2 * nx)):
            for i, x in enumerate(result["x"]):
                for sense in ("min", "max"):
                    case = result["case_" + sense][offset + i]
                    cases.append(
                        [
                            x,
                            result["sides"][i],
                            effect,
                            sense,
                            result[sense][effect][i],
                            case["case"],
                            (
                                "spans/travées "
                                + ", ".join(map(str, case["spans"]))
                                + f" · s = {case['loaded_length']:.2f} m"
                                + f" · p = {case['intensity']:.3f} kPa"
                                if case["case"] == "pedestrian"
                                else "–".join(map(str, case["axles"]))
                            ),
                            case["factor"],
                            result["model"]["live"].get("factor", 1),
                            result["model"]["live"].get("axle_factor", 1),
                            case["position"],
                            case["direction"],
                            case.get("reduction", 1.0),
                            case.get("second_position"),
                            case.get("gap"),
                        ]
                    )
        for column, label in enumerate(
            ["Réduction / Reduction", "x₂ (m)", "Entre camions / Clear gap (m)"], 13
        ):
            cases.cell(1, column, label)
        # Range columns preserve the original station/reaction column positions.
        for column, (key, unit) in enumerate(
            (("V", "kN"), ("M", "kN·m"), ("D", "mm")), 14
        ):
            ws.cell(1, column, f"Δ{key} ({unit})")
            for row, station in enumerate(result["table"], 2):
                ws.cell(row, column, station[f"{key}_max"] - station[f"{key}_min"])
    modes = wb.create_sheet("Modes")
    modes.append(
        [
            "Mode",
            "f (Hz)",
            "T (s)",
            "ω (rad/s)",
            "Masse modale" if language == "fr" else "Modal mass",
            "Cumul" if language == "fr" else "Cumulative",
            "Symétrie" if language == "fr" else "Symmetry",
        ]
    )
    try:
        modal = analyse_modal(model)
    except ValueError as error:
        if str(error) != "modal.no_mass":
            raise
        modal = None
        modes.append(
            [
                "Aucune masse : modes indisponibles."
                if language == "fr"
                else "No mass: modes unavailable."
            ]
        )
    if modal:
        cumulative = 0.0
        for mode in modal["modes"]:
            cumulative += mode["mass_ratio"]
            modes.append(
                [
                    mode["n"],
                    mode["f"],
                    mode["T"],
                    mode["omega"],
                    mode["mass_ratio"],
                    cumulative,
                    mode["symmetry"],
                ]
            )
        shapes = wb.create_sheet(
            "Formes modales" if language == "fr" else "Mode shapes"
        )
        shapes.append(
            ["x (m)"] + [f"Mode {m['n']} (|φ|max = 1)" for m in modal["modes"]]
        )
        for i, x in enumerate(modal["x"]):
            shapes.append([x] + [shape[i] for shape in modal["shapes"]])
        chart = ScatterChart()
        chart.title = "Formes modales" if language == "fr" else "Mode shapes"
        chart.x_axis.title, chart.y_axis.title = "x (m)", "φ (|φ|max = 1)"
        chart.width, chart.height = 24, 12
        for column in range(2, min(shapes.max_column, 7) + 1):
            series = Series(
                Reference(shapes, min_col=column, min_row=2, max_row=shapes.max_row),
                Reference(shapes, min_col=1, min_row=2, max_row=shapes.max_row),
                title=f"Mode {column-1}",
            )
            chart.series.append(series)
        shapes.add_chart(chart, "O2")
    if model.distribution.enabled:
        distribution_sheet(wb, model, language)
    meta = wb.create_sheet("Modèle" if language == "fr" else "Model")
    meta.append(["QuickerBridge", APP_VERSION])
    # v0.9.96: model name and export date (also in the file name).
    info = info or {}
    fr = language == "fr"
    meta.append(["Nom du modèle" if fr else "Model name", info.get("name") or ""])
    meta.append(["Date d’export" if fr else "Export date", info.get("exported") or ""])
    meta.append(["Warning / Avertissement", WARNING[language]])
    meta.append(["Version", f"QuickerBridge {APP_VERSION} · {RELEASE_DATE} · {AUTHOR}"])
    if thermal:
        thermal_input = result["model"]["thermal"]
        imposed = thermal_input.get("imposed", "thermal")
        if imposed == "thermal":
            meta.append(
                [
                    "Load case / Cas",
                    "Thermal gradient only / Gradient thermique seulement",
                ]
            )
            meta.append(["ΔT = Ttop − Tbottom (°C)", thermal_input["delta_T"]])
            meta.append(["α (10⁻⁶/°C)", thermal_input["alpha_micro"]])
            meta.append(
                ["Thermal depth / Hauteur thermique (mm)", thermal_input["depth"]]
            )
        else:
            meta.append(
                [
                    "Load case / Cas",
                    "Slab shrinkage only / Retrait de la dalle seulement"
                    if imposed == "shrinkage"
                    else "Slab creep only / Fluage de la dalle seulement",
                ]
            )
            if imposed == "creep":
                meta.append(["φ", thermal_input["creep_phi"]])
                meta.append(["σc (MPa)", thermal_input["creep_stress"]])
            meta.append(
                [
                    "Slab shortening / Raccourcissement de la dalle (10⁻⁶)",
                    result["meta"].get("slab_strain"),
                ]
            )
            meta.append(
                [
                    "Long-term section / Section à long terme",
                    f"{thermal_input['modular_factor']:g}n",
                ]
            )
        meta.append(
            ["Imposed curvature / Courbure imposée (1/m)", result["meta"]["curvature"]]
        )
    else:
        vehicle = result["model"]["live"]["vehicle"]
        dynamic_reference = (
            "CAN/CSA S6-25 · 3.8.4.5.3"
            if vehicle in ("CL625", "CL750QC")
            else "AASHTO LRFD · 33% on axles"
            if vehicle in ("HL93Truck", "HL93Tandem")
            else "None / Aucun"
            if vehicle in ("Cooper", "Maintenance")
            else "User selection / Choix utilisateur"
        )
        meta.append(["Dynamic allowance / CMD", dynamic_reference])
        meta.append(
            [
                "Live factors / Facteurs de surcharge",
                f"load/charge = {result['model']['live'].get('factor', 1)}; "
                f"axle/essieu = {result['model']['live'].get('axle_factor', 1)}",
            ]
        )
        meta.append(
            [
                "Dead factors / Facteurs permanents",
                "; ".join(
                    f"{load['name']} = {load.get('factor', 1)}"
                    for load in result["model"]["dead"]
                ),
            ]
        )
        self_weight = result["model"].get("self_weight", {})
        meta.append(
            [
                "Girder self-weight / Poids propre des poutres",
                (
                    f"applied/appliqué · steel/acier 77 kN/m³ +{self_weight['steel_increase']} % · "
                    f"NEBT +{self_weight['nebt_increase']} % · factor/facteur = {self_weight['factor']}"
                    if self_weight.get("apply")
                    else "not applied / non appliqué"
                ),
            ]
        )
    if result.get("ft"):
        ft = result["ft"]
        meta.append(
            [
                "FT S6-25 (par zone / per zone)",
                f"{ft['girder']} · {ft['state']} · "
                + "; ".join(
                    f"{z['x0']:.2f}–{z['x1']:.2f} m M×{z['FT_M']:.3f} V×{z['FT_V']:.3f}"
                    for z in ft["zones"]
                )
                + " · whole live load (trucks and lane) / toute la surcharge (camions et voie): "
                "V × FT shear, M and δ × FT moment; R: one lane / une voie; "
                "exterior girder: dead-load V × Fs / poutre ext. : V permanent × Fs",
            ]
        )
    simple = [
        str(i + 1) for i, s in enumerate(result["model"]["spans"]) if s.get("simple")
    ]
    if simple:
        meta.append(
            [
                "Isostatic spans / Travées isostatiques",
                ", ".join(simple)
                + " (moment releases at both ends / rotules aux deux extrémités)",
            ]
        )
    meta.append(
        [
            "Scope / Portée",
            "1-D Euler–Bernoulli; one lane / une voie; gross homogeneous I-section / section en I homogène brute",
        ]
    )
    if not single:
        meta.append(
            [
                "Factors / Facteurs",
                "Code dynamic/lane factors and the user load/axle factors are applied separately.",
            ]
        )
        extent = result["model"]["live"].get("lane_extent", "spans")
        meta.append(
            [
                "Lane / Voie",
                {
                    "full": "Companion UDL over the full bridge length / "
                    "charge répartie sur tout le pont.",
                    "spans": "Companion UDL on the spans that increase each effect "
                    "(S6 C3.8.4.1) / charge répartie sur les travées qui augmentent "
                    "chaque effet.",
                    "influence": "Companion UDL on the parts of each influence line "
                    "that increase the effect (S6 C3.8.4.1) / charge répartie sur les "
                    "parties de la ligne d’influence qui augmentent l’effet.",
                }[extent],
            ]
        )
        if result["model"]["live"].get("source") == "pedestrian":
            ped = result["model"].get("pedestrian", {})
            width = (result.get("vehicle") or {}).get("pedestrian_width")
            meta.append(
                [
                    "Pedestrian load / Charge piétonnière (S6-25 3.8.9)",
                    "p = 4.25 (0.5 + √(5/s)) ≤ 4.25 kPa; "
                    f"width / largeur = {width if width is None else round(width, 3)} m; "
                    "every combination of loaded spans / toutes les combinaisons "
                    "de travées chargées"
                    + (
                        "; or maintenance vehicle, not concomitant / ou véhicule "
                        "d’entretien, non concomitant"
                        if ped.get("maintenance")
                        else ""
                    ),
                ]
            )
        live = result["model"]["live"]
        if vehicle == "CL625" and live.get("evaluation", "design") != "design":
            level = live["evaluation"]
            w = 625
            q = {"A": 9, "B": 8, "CD": 7}[live.get("road_class", "A")]
            meta.append(
                [
                    "Evaluation / Évaluation (S6-25 14.9.1)",
                    f"Level / Niveau {level}: CL{level}-W, W = {w} kN "
                    f"({ {'1': 5, '2': 4, '3': 3}[level]} axles / essieux); "
                    f"lane / voie: 80 % axles + q = {q} kN/m "
                    f"(road class / classe de route "
                    f"{live.get('road_class', 'A').replace('CD', 'C/D')}).",
                ]
            )
        if (
            vehicle == "CL750QC"
            and model.live.mtq_auto
            and (live.get("evaluation", "design") == "design")
        ):
            meta.append(
                [
                    "MTQ automatic fraction / Fraction MTQ automatique",
                    "CL-750-QC lane case, MTQ A2023-05: axles at 63 % for M+ outside the M− zones of the "
                    "supports, V, single span and reactions without deck continuity; 80 % for M−, M+ in the "
                    "M− zones, reactions at continuous piers and deflections.",
                ]
            )
        if vehicle in ("HL93Truck", "HL93Tandem") and model.live.two_trucks:
            meta.append(
                [
                    "HL-93 · 90% · two trucks / deux camions",
                    "Supplementary lane case for M− around interior piers and interior vertical R only; "
                    "90% trucks + full-deck lane; 14 ft axle spacings; ≥50 ft clear headway; "
                    "truck centres in adjacent spans. FHWA-HIF-16-002 Vol.20 §6.2.1.",
                ]
            )
        meta.append(
            [
                "Envelope / Enveloppe",
                "Each station/effect can have a different governing arrangement. Reactions are reported once per support.",
            ]
        )
    meta.append(["PyCBA", result["meta"]["pycba"]])
    if not thermal:
        meta.append(["Travel step / Pas (m)", result["meta"]["travel_step"]])
    meta.append(["Masse / Mass", model.modal.mass_source])
    if modal:
        meta.append(["Masse totale / Total mass (t)", modal["mass"]["total_t"]])
        meta.append(["Masse moyenne / Mean mass (t/m)", modal["mass"]["mean_t_per_m"]])
    meta.append(["Model JSON", json.dumps(model.model_dump(), ensure_ascii=False)])
    for index, sheet in enumerate(wb):
        sheet.freeze_panes = (
            "D2"
            if sheet is ws
            else "B2"
            if sheet.title in ("Mode shapes", "Formes modales")
            else "A2"
        )
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.tabColor = "008378" if sheet is modes else "102D41"
        if sheet is not meta and sheet.max_row > 1 and (sheet is not modes or modal):
            table = Table(displayName=f"QBTable{index+1}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            sheet.add_table(table)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="102D41")
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(
                wrap_text=True, vertical="center", horizontal="center"
            )
            cell.border = Border(right=Side(style="thin", color="FFFFFF"))
        sheet.row_dimensions[1].height = 42
        for column in sheet.columns:
            header = str(column[0].value or "")
            sheet.column_dimensions[column[0].column_letter].width = min(
                32, max(14, len(header) * 0.7)
            )
            for cell in column[1:]:
                cell.font = Font(name="Arial", size=10, color="19384B")
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="right"
                    if isinstance(cell.value, (float, int))
                    else "left",
                )
                if isinstance(cell.value, (float, int)):
                    cell.number_format = "#,##0.000"
                    if header in (
                        "Mode",
                        "Span",
                        "Travée",
                        "Station",
                        "Support",
                        "Appui",
                    ):
                        cell.number_format = "0"
                    elif "(kN" in header or header.startswith("Δ"):
                        cell.number_format = "#,##0.00"
                    if sheet is modes and cell.column in (5, 6):
                        cell.number_format = "0.0%"
                if isinstance(cell.value, str) and cell.value.startswith(
                    ("=", "+", "-", "@")
                ):
                    cell.value = "'" + cell.value
        sheet.sheet_view.showGridLines = False
    meta.column_dimensions["A"].width = 36
    meta.column_dimensions["B"].width = 96
    for row in meta.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        text = str(row[1].value or "")
        meta.row_dimensions[row[0].row].height = min(
            60, 16 * max(1, (len(text) + 90) // 91)
        )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
